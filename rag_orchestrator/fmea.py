"""Evidence-bound FMEA task generation and review.

The generator is intentionally deterministic: it transforms approved graph
statements into an auditable candidate and never invents risk scores or missing
professional facts.  Missing and conflicting fields stay visible for a human
reviewer, as required by the delivery workflow.
"""
# ruff: noqa: TRY003

from __future__ import annotations

import csv
import io
import json
from collections.abc import Mapping, Sequence
from dataclasses import replace
from typing import Any

from core_domain.delivery import (
    FMEA_FIELDS,
    FMEAItem,
    FMEATaskRequest,
    FMEATaskResult,
    GraphStatement,
    IssueSeverity,
    QualityIssue,
    ReviewDecision,
    TaskError,
    TaskStatus,
)
from rag_orchestrator.fmea_templates import FMEATemplate, FMEATemplateError, FMEATemplateRegistry
from storage_layer.governance_store import GovernanceError, GovernanceStore

FMEA_DOCX_HEADER_LABELS = {
    "item_id": "项目 ID",
    "equipment": "设备",
    "component": "部件",
    "failure_mode": "故障模式",
    "cause": "原因",
    "effect": "后果",
    "detection_method": "检测方法",
    "recommended_action": "建议措施",
}
FMEA_DOCX_COLUMN_WIDTHS = (0.75, 1.05, 1.20, 1.20, 1.15, 1.30, 1.45, 1.70)


class FMEAService:
    """Run, review, publish, and export evidence-grounded FMEA tasks."""

    def __init__(self, store: GovernanceStore, *, template_registry: FMEATemplateRegistry | None = None) -> None:
        self.store = store
        self.template_registry = template_registry or FMEATemplateRegistry()

    def run(self, request: FMEATaskRequest) -> FMEATaskResult:
        template = self.template_registry.get(request.template, request.template_version)
        governed_request = replace(
            request,
            template=template.template_id,
            template_version=template.version,
            metadata={**request.metadata, "template_lineage": template.lineage()},
        )
        task = self.store.create_fmea_task(governed_request)
        self.store.save_fmea_result(
            task.task_id,
            status=TaskStatus.RUNNING,
            items=(),
            reason="generation_started",
            actor=request.requested_by,
        )
        try:
            graph = self.store.get_graph_version(governed_request.graph_version_id)
            items = build_fmea_items(graph.statements)
            if not items:
                return self.store.save_fmea_result(
                    task.task_id,
                    status=TaskStatus.FAILED,
                    items=(),
                    errors=(
                        TaskError(
                            code="fmea_failure_mode_missing",
                            message="No HAS_FAILURE_MODE statements were available for the selected graph version.",
                            details={"graph_version_id": governed_request.graph_version_id},
                        ),
                    ),
                    reason="generation_failed",
                    actor=request.requested_by,
                )
            return self.store.save_fmea_result(
                task.task_id,
                status=TaskStatus.NEEDS_REVIEW,
                items=items,
                reason="candidate_generated",
                actor=request.requested_by,
            )
        except Exception as exc:
            self.store.save_fmea_result(
                task.task_id,
                status=TaskStatus.FAILED,
                items=(),
                errors=(
                    TaskError(
                        code=str(getattr(exc, "code", "fmea_generation_failed")),
                        message=str(exc),
                        details=dict(getattr(exc, "details", {}) or {}),
                    ),
                ),
                reason="generation_failed",
                actor=request.requested_by,
            )
            raise

    def review(
        self,
        task_id: str,
        *,
        reviewer: str,
        decision: ReviewDecision | str,
        comment: str = "",
        corrections: Mapping[str, Any] | None = None,
        review_metadata: Mapping[str, Any] | None = None,
    ) -> FMEATaskResult:
        decision = decision if isinstance(decision, ReviewDecision) else ReviewDecision(str(decision))
        task = self.store.get_fmea_task(task_id)
        if task.status not in {TaskStatus.NEEDS_REVIEW, TaskStatus.APPROVED}:
            raise GovernanceError(f"Task {task_id} is not reviewable from status {task.status.value}")

        self._validate_correction_evidence(task, corrections or {})
        corrected_items = _apply_corrections(task.items, corrections or {})
        if decision in {ReviewDecision.APPROVE, ReviewDecision.CONFIRM}:
            corrected_items = tuple(replace(item, review_status="approved") for item in corrected_items)
            status = TaskStatus.APPROVED
        elif decision is ReviewDecision.REJECT:
            corrected_items = tuple(replace(item, review_status="rejected") for item in corrected_items)
            status = TaskStatus.NEEDS_REVIEW
        else:
            status = TaskStatus.NEEDS_REVIEW

        recorded_corrections: dict[str, Any] = dict(corrections or {})
        if review_metadata:
            if "items" not in recorded_corrections:
                recorded_corrections = {"items": recorded_corrections}
            recorded_corrections["_review"] = dict(review_metadata)
        self.store.record_review(
            target_type="fmea",
            target_id=task_id,
            reviewer=reviewer,
            decision=decision,
            comment=comment,
            corrections=recorded_corrections,
        )
        return self.store.save_fmea_result(
            task_id,
            status=status,
            items=corrected_items,
            errors=task.errors,
            reason=f"review_{decision.value}",
            actor=reviewer,
        )

    def publish(self, task_id: str) -> FMEATaskResult:
        task = self.store.get_fmea_task(task_id)
        if task.status is not TaskStatus.APPROVED:
            raise GovernanceError(
                "FMEA task must be approved before publication",
                code="fmea_not_approved",
                details={"status": task.status.value},
            )
        if task.request.metadata.get("require_field_review"):
            reviews = self.store.list_reviews("fmea", task_id)
            review_metadata = dict(reviews[-1].corrections.get("_review") or {}) if reviews else {}
            decisions = {
                str(item_id): {str(field): str(decision) for field, decision in dict(fields).items()}
                for item_id, fields in dict(review_metadata.get("field_decisions") or {}).items()
            }
            expected = {
                item.item_id: {field for field, value in item.fields.items() if value not in (None, "")}
                for item in task.items
            }
            complete = (
                review_metadata.get("review_mode") == "field_by_field"
                and set(decisions) == set(expected)
                and all(set(decisions[item_id]) == fields for item_id, fields in expected.items())
                and all(
                    decision == "approve"
                    for item_decisions in decisions.values()
                    for decision in item_decisions.values()
                )
            )
            if not complete:
                raise GovernanceError(
                    "FMEA publication requires an explicit approval for every populated professional field",
                    code="fmea_field_review_incomplete",
                    details={
                        "expected_field_count": sum(len(fields) for fields in expected.values()),
                        "approved_field_count": sum(
                            decision == "approve"
                            for item_decisions in decisions.values()
                            for decision in item_decisions.values()
                        ),
                    },
                )
        template = self.template_registry.get(task.request.template, task.request.template_version)
        self._validate_publishable_evidence(task, template)
        return self.store.publish_fmea_task(task_id)

    def result_payload(self, task_id: str) -> dict[str, Any]:
        """Return a self-contained task result with resolved source locators and reviews."""

        task = self.store.get_fmea_task(task_id)
        payload = task.to_dict()
        payload["template_definition"] = self.template_registry.get(
            task.request.template, task.request.template_version
        ).to_dict()
        payload["reviews"] = [item.to_dict() for item in self.store.list_reviews("fmea", task_id)]
        for item, item_payload in zip(task.items, payload["items"], strict=True):
            item_payload["field_evidence_details"] = self._resolved_field_evidence(task, item)
        return payload

    def export_json(self, task_id: str) -> str:
        self._published(task_id)
        return json.dumps(self.result_payload(task_id), ensure_ascii=False, indent=2)

    def export_csv(self, task_id: str) -> str:
        task = self._published(task_id)
        payload = self.result_payload(task_id)
        output = io.StringIO(newline="")
        fieldnames = [
            "task_id",
            "task_status",
            "graph_version_id",
            "document_version_ids",
            "template",
            "template_version",
            "item_id",
            *FMEA_FIELDS,
            *[f"{field}_evidence" for field in FMEA_FIELDS],
            *[f"{field}_evidence_details" for field in FMEA_FIELDS],
            "issues",
            "review_status",
            "review_history",
            "state_history",
        ]
        writer = csv.DictWriter(output, fieldnames=fieldnames)
        writer.writeheader()
        for row in self._csv_rows(task, payload):
            writer.writerow(row)
        return output.getvalue()

    def export_docx(self, task_id: str) -> bytes:  # noqa: C901
        """Create a landscape, evidence-bearing formal FMEA report."""

        task = self._published(task_id)
        payload = self.result_payload(task_id)
        template = self.template_registry.get(task.request.template, task.request.template_version)
        report_layout = template.report_layout
        try:
            from docx import Document
            from docx.enum.section import WD_ORIENT
            from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.oxml.ns import qn
            from docx.shared import Inches, Pt, RGBColor
        except ModuleNotFoundError as exc:  # pragma: no cover - dependency gate
            raise FMEATemplateError("python-docx is required for formal FMEA DOCX export") from exc

        document = Document()
        section = document.sections[0]
        if report_layout.get("orientation") == "landscape":
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width, section.page_height = section.page_height, section.page_width
        margin = max(0.2, min(float(report_layout.get("margin_inches") or 0.45), 1.5))
        section.top_margin = Inches(margin)
        section.bottom_margin = Inches(margin)
        section.left_margin = Inches(margin)
        section.right_margin = Inches(margin)
        normal = document.styles["Normal"]
        font_name = str(report_layout.get("font") or "Microsoft YaHei")
        normal.font.name = font_name
        normal.font.size = Pt(8)
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)

        title = document.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run(str(report_layout.get("title") or "PowerRAG 可信 FMEA 正式交付表"))
        run.bold = True
        run.font.size = Pt(18)
        accent = str(report_layout.get("accent_color") or "1F4E79").lstrip("#")
        if len(accent) != 6 or any(character not in "0123456789abcdefABCDEF" for character in accent):
            accent = "1F4E79"
        run.font.color.rgb = RGBColor.from_string(accent.upper())
        metadata_table = document.add_table(rows=3, cols=4)
        metadata_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        metadata_table.style = "Table Grid"
        metadata = (
            ("任务 ID", task.task_id, "状态", task.status.value),
            ("图谱版本", task.request.graph_version_id, "模板", f"{task.request.template}@{task.request.template_version or ''}"),
            ("资料版本", "\n".join(task.request.document_version_ids), "项目 ID", task.request.project_id),
        )
        for row, values in zip(metadata_table.rows, metadata, strict=True):
            for cell, value in zip(row.cells, values, strict=True):
                cell.text = str(value)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        document.add_paragraph()
        headers = ("item_id", *FMEA_FIELDS)
        table = document.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        for index, (cell, header) in enumerate(zip(table.rows[0].cells, headers, strict=True)):
            cell.width = Inches(FMEA_DOCX_COLUMN_WIDTHS[index])
            table.columns[index].width = Inches(FMEA_DOCX_COLUMN_WIDTHS[index])
            cell.text = FMEA_DOCX_HEADER_LABELS[header]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for header_run in paragraph.runs:
                    header_run.bold = True
                    header_run.font.color.rgb = RGBColor(255, 255, 255)
            cell._tc.get_or_add_tcPr().append(_docx_cell_shading(accent.upper()))
        for item in payload.get("items") or []:
            cells = table.add_row().cells
            values = {"item_id": item.get("item_id"), **dict(item.get("fields") or {})}
            evidence = dict(item.get("field_evidence") or {})
            for index, (cell, field_name) in enumerate(zip(cells, headers, strict=True)):
                cell.width = Inches(FMEA_DOCX_COLUMN_WIDTHS[index])
                value = str(values.get(field_name) or "")
                evidence_ids = evidence.get(field_name) or []
                cell.text = value
                if field_name != "item_id" and evidence_ids:
                    evidence_paragraph = cell.add_paragraph(f"证据: {', '.join(evidence_ids)}")
                    for evidence_run in evidence_paragraph.runs:
                        evidence_run.font.size = Pt(6.5)
                        evidence_run.font.color.rgb = RGBColor(89, 89, 89)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

        document.add_paragraph("说明: 每个非空专业字段均须保留独立证据 ID; 无批准评分政策时 S/O/D 与 RPN 保持为空。")
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()

    def verify_export_consistency(self, task_id: str) -> dict[str, Any]:  # noqa: C901
        """Round-trip JSON and CSV exports and compare their governed fields."""

        task = self._published(task_id)
        json_payload = json.loads(self.export_json(task_id))
        csv_rows = list(csv.DictReader(io.StringIO(self.export_csv(task_id))))
        docx_rows = _read_fmea_docx_rows(self.export_docx(task_id))
        mismatches: list[dict[str, Any]] = []
        json_items = list(json_payload.get("items") or [])
        expected_rows = self._csv_rows(task, json_payload)
        if len(json_items) != len(csv_rows):
            mismatches.append({
                "code": "row_count_mismatch",
                "json_rows": len(json_items),
                "csv_rows": len(csv_rows),
            })
        if len(json_items) != len(docx_rows):
            mismatches.append({
                "code": "docx_row_count_mismatch",
                "json_rows": len(json_items),
                "docx_rows": len(docx_rows),
            })
        for index, expected in enumerate(expected_rows):
            if index >= len(csv_rows):
                break
            row = csv_rows[index]
            for column, expected_value in expected.items():
                actual_value = str(row.get(column) or "")
                if expected_value != actual_value:
                    mismatches.append({
                        "code": "column_mismatch",
                        "row": index + 1,
                        "column": column,
                        "json": expected_value,
                        "csv": actual_value,
                    })
            if index < len(docx_rows):
                docx_row = docx_rows[index]
                for field_name in ("item_id", *FMEA_FIELDS):
                    expected_value = str(json_items[index].get("fields", {}).get(field_name) or "")
                    if field_name == "item_id":
                        expected_value = str(json_items[index].get("item_id") or "")
                    if docx_row.get(field_name, "") != expected_value:
                        mismatches.append({
                            "code": "docx_field_mismatch",
                            "row": index + 1,
                            "field": field_name,
                            "json": expected_value,
                            "docx": docx_row.get(field_name, ""),
                        })
        return {
            "task_id": task.task_id,
            "status": task.status.value,
            "consistent": not mismatches,
            "json_rows": len(json_items),
            "csv_rows": len(csv_rows),
            "docx_rows": len(docx_rows),
            "mismatches": mismatches,
        }

    def _csv_rows(self, task: FMEATaskResult, payload: Mapping[str, Any]) -> list[dict[str, str]]:
        reviews = json.dumps(payload.get("reviews") or [], ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        state_history = json.dumps(
            payload.get("state_history") or [], ensure_ascii=False, sort_keys=True, separators=(",", ":")
        )
        rows: list[dict[str, str]] = []
        for item in payload.get("items") or []:
            fields = dict(item.get("fields") or {})
            field_evidence = dict(item.get("field_evidence") or {})
            details = dict(item.get("field_evidence_details") or {})
            row = {
                "task_id": task.task_id,
                "task_status": task.status.value,
                "graph_version_id": task.request.graph_version_id,
                "document_version_ids": "|".join(task.request.document_version_ids),
                "template": task.request.template,
                "template_version": task.request.template_version or "",
                "item_id": str(item.get("item_id") or ""),
                "issues": json.dumps(
                    item.get("issues") or [], ensure_ascii=False, sort_keys=True, separators=(",", ":")
                ),
                "review_status": str(item.get("review_status") or ""),
                "review_history": reviews,
                "state_history": state_history,
            }
            row.update({field: str(fields.get(field) or "") for field in FMEA_FIELDS})
            row.update({
                f"{field}_evidence": "|".join(str(value) for value in field_evidence.get(field) or ())
                for field in FMEA_FIELDS
            })
            row.update({
                f"{field}_evidence_details": json.dumps(
                    details.get(field) or [], ensure_ascii=False, sort_keys=True, separators=(",", ":")
                )
                for field in FMEA_FIELDS
            })
            rows.append(row)
        return rows

    def _resolved_field_evidence(self, task: FMEATaskResult, item: FMEAItem) -> dict[str, list[dict[str, Any]]]:
        allowed_documents = set(task.request.document_version_ids)
        resolved: dict[str, list[dict[str, Any]]] = {}
        for field_name in FMEA_FIELDS:
            locators = []
            for evidence_id in item.field_evidence.get(field_name, ()):
                locator = self.store.get_evidence(evidence_id)
                if locator.document_version_id not in allowed_documents:
                    raise GovernanceError(
                        f"FMEA evidence is outside the selected document versions: {evidence_id}",
                        code="fmea_evidence_out_of_scope",
                        details={"evidence_id": evidence_id, "field": field_name},
                    )
                locators.append(locator.to_dict())
            resolved[field_name] = locators
        return resolved

    def _validate_correction_evidence(  # noqa: C901
        self,
        task: FMEATaskResult,
        corrections: Mapping[str, Any],
    ) -> None:
        raw_items = corrections.get("items", corrections)
        if not isinstance(raw_items, Mapping):
            raise TypeError("FMEA corrections must be a mapping keyed by item_id")
        known_items = {item.item_id for item in task.items}
        unknown_items = sorted(str(item_id) for item_id in raw_items if str(item_id) not in known_items)
        if unknown_items:
            raise GovernanceError(
                f"Unknown FMEA correction items: {unknown_items}",
                code="unknown_fmea_item",
                details={"item_ids": unknown_items},
            )
        allowed_documents = set(task.request.document_version_ids)
        for item_id, item_patch in raw_items.items():
            if not isinstance(item_patch, Mapping):
                raise TypeError(f"FMEA correction for {item_id} must be a mapping")
            for field_name, raw_value in item_patch.items():
                if field_name not in FMEA_FIELDS:
                    raise ValueError(f"Unknown FMEA correction field: {field_name}")
                if not isinstance(raw_value, Mapping) or "evidence_ids" not in raw_value:
                    continue
                for evidence_id in raw_value.get("evidence_ids") or ():
                    try:
                        locator = self.store.get_evidence(str(evidence_id))
                    except GovernanceError as exc:
                        raise GovernanceError(
                            f"Unknown evidence in FMEA correction: {evidence_id}",
                            code="unknown_fmea_evidence",
                            details={"item_id": item_id, "field": field_name, "evidence_id": evidence_id},
                        ) from exc
                    if locator.document_version_id not in allowed_documents:
                        raise GovernanceError(
                            f"FMEA correction evidence is outside selected documents: {evidence_id}",
                            code="fmea_evidence_out_of_scope",
                            details={"item_id": item_id, "field": field_name, "evidence_id": evidence_id},
                        )

    def _validate_publishable_evidence(self, task: FMEATaskResult, template: FMEATemplate) -> None:
        per_field_required = bool(template.evidence_policy.get("per_field_required", False))
        for item in task.items:
            self._resolved_field_evidence(task, item)
            if not per_field_required:
                continue
            unsupported = [
                field_name
                for field_name, value in item.fields.items()
                if value and not item.field_evidence.get(field_name)
            ]
            if unsupported:
                raise GovernanceError(
                    f"FMEA item {item.item_id} has populated fields without evidence: {unsupported}",
                    code="fmea_evidence_required",
                    details={"item_id": item.item_id, "fields": unsupported},
                )

    def _published(self, task_id: str) -> FMEATaskResult:
        task = self.store.get_fmea_task(task_id)
        if task.status is not TaskStatus.PUBLISHED:
            raise GovernanceError(
                "Only published FMEA tasks can be exported",
                code="fmea_not_published",
                details={"status": task.status.value},
            )
        return task


def _docx_cell_shading(fill: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    return shading


def _read_fmea_docx_rows(content: bytes) -> list[dict[str, str]]:
    from docx import Document

    document = Document(io.BytesIO(content))
    expected_headers = ["item_id", *FMEA_FIELDS]
    expected_labels = [FMEA_DOCX_HEADER_LABELS[header] for header in expected_headers]
    for table in document.tables:
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        if headers != expected_labels:
            continue
        return [
            {
                header: row.cells[index].paragraphs[0].text.strip()
                for index, header in enumerate(expected_headers)
            }
            for row in table.rows[1:]
        ]
    return []


def build_fmea_items(statements: Sequence[GraphStatement]) -> tuple[FMEAItem, ...]:
    """Build the minimum gas-turbine FMEA template from governed graph facts."""
    statements = tuple(statements)
    node_types = _node_types(statements)
    parent_edges = [item for item in statements if item.predicate == "PART_OF"]
    failure_edges = [item for item in statements if item.predicate == "HAS_FAILURE_MODE"]
    by_subject_relation: dict[tuple[str, str], list[GraphStatement]] = {}
    for statement in statements:
        by_subject_relation.setdefault((statement.subject, statement.predicate), []).append(statement)

    items: list[FMEAItem] = []
    for index, failure_edge in enumerate(failure_edges, start=1):
        component = failure_edge.subject
        failure_mode = failure_edge.object_name
        equipment, equipment_edges = _resolve_equipment(component, parent_edges, node_types)
        if failure_edge.subject_type == "EQUIPMENT":
            equipment = component
            component_value: str | None = None
        else:
            component_value = component

        cause_edges = by_subject_relation.get((failure_mode, "CAUSED_BY"), [])
        effect_edges = by_subject_relation.get((failure_mode, "HAS_EFFECT"), [])
        detection_edges = by_subject_relation.get((failure_mode, "DETECTED_BY"), [])
        action_edges = by_subject_relation.get((failure_mode, "MITIGATED_BY"), [])

        fields: dict[str, str | None] = {
            "equipment": equipment,
            "component": component_value,
            "failure_mode": failure_mode,
            "cause": _join_objects(cause_edges),
            "effect": _join_objects(effect_edges),
            "detection_method": _join_objects(detection_edges),
            "recommended_action": _join_objects(action_edges),
        }
        field_evidence: dict[str, tuple[str, ...]] = {
            "equipment": _edge_evidence(equipment_edges or [failure_edge]),
            "component": failure_edge.evidence_ids,
            "failure_mode": failure_edge.evidence_ids,
            "cause": _edge_evidence(cause_edges),
            "effect": _edge_evidence(effect_edges),
            "detection_method": _edge_evidence(detection_edges),
            "recommended_action": _edge_evidence(action_edges),
        }
        issues = _fmea_issues(
            item_id=f"FMEA-{index:04d}",
            fields=fields,
            field_evidence=field_evidence,
            field_edges={
                "cause": cause_edges,
                "effect": effect_edges,
                "detection_method": detection_edges,
                "recommended_action": action_edges,
            },
        )
        source_statement_ids = tuple(
            dict.fromkeys(
                statement.statement_id
                for statement in [
                    failure_edge,
                    *equipment_edges,
                    *cause_edges,
                    *effect_edges,
                    *detection_edges,
                    *action_edges,
                ]
            )
        )
        items.append(
            FMEAItem(
                item_id=f"FMEA-{index:04d}",
                fields=fields,
                field_evidence=field_evidence,
                issues=issues,
                metadata={"source_statement_ids": list(source_statement_ids)},
            )
        )
    return tuple(items)


def _resolve_equipment(
    component: str,
    parent_edges: Sequence[GraphStatement],
    node_types: Mapping[str, str],
) -> tuple[str | None, list[GraphStatement]]:
    by_child: dict[str, list[GraphStatement]] = {}
    for edge in parent_edges:
        by_child.setdefault(edge.subject, []).append(edge)
    current = component
    visited = {current}
    path: list[GraphStatement] = []
    while by_child.get(current):
        edge = by_child[current][0]
        path.append(edge)
        current = edge.object_name
        if current in visited:
            break
        visited.add(current)
        if node_types.get(current) == "EQUIPMENT" or edge.object_type == "EQUIPMENT":
            return current, path
    return (current if path else None), path


def _node_types(statements: Sequence[GraphStatement]) -> dict[str, str]:
    types: dict[str, str] = {}
    for item in statements:
        types.setdefault(item.subject, item.subject_type)
        types.setdefault(item.object_name, item.object_type)
    return types


def _join_objects(edges: Sequence[GraphStatement]) -> str | None:
    values = list(dict.fromkeys(edge.object_name for edge in edges if edge.object_name))
    return " / ".join(values) if values else None


def _edge_evidence(edges: Sequence[GraphStatement]) -> tuple[str, ...]:
    return tuple(dict.fromkeys(evidence_id for edge in edges for evidence_id in edge.evidence_ids))


def _fmea_issues(
    *,
    item_id: str,
    fields: Mapping[str, str | None],
    field_evidence: Mapping[str, tuple[str, ...]],
    field_edges: Mapping[str, Sequence[GraphStatement]],
) -> tuple[QualityIssue, ...]:
    specs: list[tuple[str, str, IssueSeverity, tuple[str, ...], dict[str, Any]]] = []
    for field_name in FMEA_FIELDS:
        value = fields.get(field_name)
        citations = field_evidence.get(field_name, ())
        if not value:
            specs.append((
                "missing_field",
                f"FMEA field {field_name} is unknown and requires human review.",
                IssueSeverity.WARNING,
                (),
                {"field": field_name},
            ))
        elif not citations:
            specs.append((
                "insufficient_evidence",
                f"FMEA field {field_name} has a value but no source evidence.",
                IssueSeverity.ERROR,
                (),
                {"field": field_name},
            ))
    for field_name, edges in field_edges.items():
        values = {edge.object_name for edge in edges if edge.object_name}
        if len(values) > 1:
            specs.append((
                "field_conflict",
                f"FMEA field {field_name} has multiple source values: {', '.join(sorted(values))}.",
                IssueSeverity.WARNING,
                _edge_evidence(edges),
                {"field": field_name, "values": sorted(values)},
            ))
    return tuple(
        QualityIssue(
            issue_id=f"{item_id}:Q{index:03d}",
            code=code,
            message=message,
            severity=severity,
            evidence_ids=evidence_ids,
            metadata=metadata,
        )
        for index, (code, message, severity, evidence_ids, metadata) in enumerate(specs, start=1)
    )


def _apply_corrections(items: Sequence[FMEAItem], corrections: Mapping[str, Any]) -> tuple[FMEAItem, ...]:
    raw_items = corrections.get("items", corrections)
    if not isinstance(raw_items, Mapping):
        raise TypeError("FMEA corrections must be a mapping keyed by item_id")
    corrected: list[FMEAItem] = []
    for item in items:
        item_patch = raw_items.get(item.item_id, {})
        if not isinstance(item_patch, Mapping) or not item_patch:
            corrected.append(item)
            continue
        fields = dict(item.fields)
        field_evidence = dict(item.field_evidence)
        changed_fields: list[str] = []
        for field_name, raw_value in item_patch.items():
            if field_name not in FMEA_FIELDS:
                raise ValueError(f"Unknown FMEA correction field: {field_name}")
            if isinstance(raw_value, Mapping):
                value = raw_value.get("value")
                citations = tuple(
                    str(value).strip() for value in raw_value.get("evidence_ids", ()) if str(value).strip()
                )
            else:
                value = raw_value
                citations = field_evidence.get(field_name, ())
            fields[field_name] = str(value).strip() if value not in (None, "") else None
            field_evidence[field_name] = citations
            changed_fields.append(field_name)

        issues = tuple(
            replace(issue, resolved=True) if issue.metadata.get("field") in changed_fields else issue
            for issue in item.issues
        )
        for field_name in changed_fields:
            if fields.get(field_name) and not field_evidence.get(field_name):
                issues += (
                    QualityIssue(
                        issue_id=f"{item.item_id}:H{len(issues) + 1:03d}",
                        code="human_value_without_evidence",
                        message=f"Human correction for {field_name} has no evidence binding.",
                        severity=IssueSeverity.WARNING,
                        metadata={"field": field_name},
                    ),
                )
        corrected.append(
            replace(
                item,
                fields=fields,
                field_evidence=field_evidence,
                issues=issues,
                review_status="modified",
                metadata={**item.metadata, "human_modified_fields": changed_fields},
            )
        )
    return tuple(corrected)
