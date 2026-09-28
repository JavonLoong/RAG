# -*- coding: utf-8 -*-
"""在已改过的展示稿上补专业术语页脚，不重排正文。"""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt
from pptx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
SRC = HERE / "纪文龙_M2-M5工作展示_20260923.pptx"

BLACK = RGBColor(0x00, 0x00, 0x00)
FONT = "微软雅黑"
FOOT_TOP = 5989320
FOOT_H = 620000

# 每页底部注释：只解释本页出现的术语
FOOTERS = [
    "2026-09-23    FMEA：故障模式与影响分析。只出带证据的字段，不自动打风险分。",
    "FMEA：故障模式与影响分析。TECH-15 / 16：参考机性能、60MB 实传两项验收。",
    "解析  —  入库  —  图谱  —  FMEA\nS/O/D：严重度 / 发生频度 / 探测度。RPN：风险优先数，无批准政策则空着。",
    "OCR：光学字符识别。Recall@5：前 5 条检索是否命中。金标：人工标准答案。门禁：验收门槛。",
    "AC：验收条款。E2E：端到端闭环。Recall@5：前 5 条检索命中率。TECH-15 / 16：性能与大文件实传。",
]


def _font(run, size):
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    run.font.bold = False
    run.font.name = FONT
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", FONT)


def _set_footer(shape, value: str, size=13):
    shape.top = Emu(FOOT_TOP)
    shape.height = Emu(FOOT_H)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, part in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(2)
        run = p.add_run()
        run.text = part
        _font(run, size)


def _add_footer(slide, value: str):
    box = slide.shapes.add_textbox(Inches(0.45), Emu(FOOT_TOP), Inches(12.4), Emu(FOOT_H))
    _set_footer(box, value)
    return box


def _is_footer(shape) -> bool:
    if not shape.has_text_frame:
        return False
    return int(shape.top) >= 5_700_000


def main():
    prs = Presentation(str(SRC))
    if len(prs.slides) != len(FOOTERS):
        raise SystemExit(f"slide count {len(prs.slides)} != {len(FOOTERS)}")
    for slide, footer in zip(prs.slides, FOOTERS):
        existing = [sh for sh in slide.shapes if _is_footer(sh)]
        if existing:
            _set_footer(existing[0], footer)
            for extra in existing[1:]:
                extra.text_frame.clear()
        else:
            _add_footer(slide, footer)
    prs.save(str(SRC))
    print(SRC)


if __name__ == "__main__":
    main()
