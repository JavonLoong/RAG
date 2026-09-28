# -*- coding: utf-8 -*-
"""全项目整合版：技术链路 + 现网全部接口。白底黑字，不加人称。"""

from __future__ import annotations

from html import escape
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
NAME = "PowerRAG_整合版_技术链路与接口_20260923"
OUT = HERE / f"{NAME}.pptx"
HTML = HERE / f"{NAME}.html"
MD = HERE / f"{NAME}.md"
USER_PPT = Path(r"D:\虚拟C盘") / f"{NAME}.pptx"
USER_HTML = Path(r"D:\虚拟C盘") / f"{NAME}.html"
USER_MD = Path(r"D:\虚拟C盘") / f"{NAME}.md"
ROUTE_FILE = Path(r"D:\虚拟C盘\PowerRAG-latest\_all_routes.txt")

BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
W = Inches(13.333)
H = Inches(7.5)
FONT = "微软雅黑"
FOOT_TOP = 5989320

DESC: dict[tuple[str, str], str] = {
    ("GET", "/"): "打开控台首页",
    ("GET", "/api/health"): "健康检查，版本号 2.1.0",
    ("GET", "/docs"): "交互式接口文档",
    ("GET", "/docs/oauth2-redirect"): "文档页登录回跳",
    ("GET", "/openapi.json"): "机器可读接口清单",
    ("GET", "/redoc"): "另一套接口文档页",
    ("POST", "/api/upload"): "上传一份资料到工作台",
    ("GET", "/api/uploads"): "列出已上传资料",
    ("DELETE", "/api/uploads/{filename}"): "删除一份上传",
    ("POST", "/api/uploads/delete"): "批量删除上传",
    ("POST", "/api/process"): "解析已上传资料并写入向量库",
    ("POST", "/api/ingest"): "按统一入库规则写入检索库",
    ("POST", "/api/public-books-json/ingest"): "从公开书目 JSON 入库",
    ("GET", "/api/stats"): "看库规模和集合统计",
    ("GET", "/api/search"): "关键词检索",
    ("POST", "/api/search"): "带条件的检索",
    ("POST", "/api/query"): "统一问答：普通检索或沿图回答",
    ("POST", "/api/benchmark"): "检索评测",
    ("DELETE", "/api/collections/{name}"): "删除一个检索集合",
    ("GET", "/api/export"): "导出全部集合",
    ("GET", "/api/export/{collection_name}"): "导出指定集合",
    ("GET", "/api/chroma/export"): "打包向量库目录",
    ("GET", "/api/logs"): "列出操作日志",
    ("GET", "/api/logs/{filename}"): "读一份日志",
    ("GET", "/api/logs/{filename}/progress"): "看处理进度",
    ("POST", "/api/retrieval/policies/propose"): "提出检索策略草案",
    ("POST", "/api/retrieval/policies/approve"): "批准检索策略",
    ("POST", "/api/retrieval/policies/reject"): "驳回检索策略",
    ("POST", "/api/retrieval/policies/promote"): "把策略推到正式",
    ("POST", "/api/retrieval/policies/rollback"): "策略退回上一版",
    ("GET", "/api/retrieval/policies/history"): "策略变更历史",
    ("POST", "/api/retrieval/policies/roles/upsert"): "写入角色权限",
    ("GET", "/api/retrieval/policies/notification-recipients"): "列出通知接收人",
    ("POST", "/api/retrieval/policies/notification-recipients/upsert"): "写入通知接收人",
    ("POST", "/api/retrieval/policies/directory/sync"): "同步人员目录",
    ("GET", "/api/retrieval/policies/notifications"): "列出策略通知",
    ("POST", "/api/retrieval/policies/notifications/dispatch"): "发出策略通知",
    ("GET", "/api/retrieval/policies/identity-provider"): "读身份源配置",
    ("POST", "/api/retrieval/policies/identity-provider/upsert"): "写身份源配置",
    ("POST", "/api/retrieval/policies/identity-provider/login-url"): "取登录地址",
    ("GET", "/api/retrieval/policies/identity-provider/callback"): "登录回跳",
    ("POST", "/api/retrieval/policies/identity-provider/token"): "换登录令牌",
    ("POST", "/api/retrieval/policies/identity-provider/session"): "建立登录会话",
    ("POST", "/api/retrieval/policies/identity-provider/session/refresh"): "刷新会话",
    ("GET", "/api/retrieval/policies/identity-provider/sessions"): "列出会话",
    ("DELETE", "/api/retrieval/policies/identity-provider/sessions/{session_id}"): "注销一个会话",
    ("POST", "/api/retrieval/policies/identity-provider/sessions/rotate-key"): "轮换会话密钥",
    ("GET", "/api/retrieval/policies/identity-provider/sessions/key-status"): "看密钥状态",
    ("POST", "/api/retrieval/policies/identity-provider/logout"): "退出登录",
    ("GET", "/api/delivery/identity"): "当前操作人",
    ("GET", "/api/delivery/projects"): "列出项目",
    ("POST", "/api/delivery/projects"): "新建项目",
    ("GET", "/api/delivery/projects/{project_id}"): "读一个项目",
    ("PATCH", "/api/delivery/projects/{project_id}"): "改项目信息",
    ("GET", "/api/delivery/projects/{project_id}/health"): "项目健康",
    ("GET", "/api/delivery/projects/{project_id}/metrics/http"): "项目接口耗时",
    ("GET", "/api/delivery/projects/{project_id}/audit"): "项目审计记录",
    ("POST", "/api/delivery/projects/{source_project_id}/copy-template"): "按模板复制项目",
    ("POST", "/api/delivery/projects/{project_id}/export-package"): "导出项目整包",
    ("POST", "/api/delivery/projects/restore"): "从整包恢复项目",
    ("GET", "/api/delivery/projects/{project_id}/providers"): "列出模型提供方",
    ("POST", "/api/delivery/projects/{project_id}/providers"): "写入模型提供方",
    ("GET", "/api/delivery/projects/{project_id}/providers/health"): "提供方是否可用",
    ("GET", "/api/delivery/projects/{project_id}/graph-schemas"): "列出图谱规则",
    ("POST", "/api/delivery/projects/{project_id}/graph-schemas"): "登记图谱规则",
    ("GET", "/api/delivery/projects/{project_id}/graph-schemas/{schema_id}"): "读一条图谱规则",
    ("POST", "/api/delivery/projects/{project_id}/graph-schemas/{schema_id}/{version}/approve"): "批准图谱规则",
    ("GET", "/api/delivery/projects/{project_id}/fmea-templates"): "列出项目出表模板",
    ("POST", "/api/delivery/projects/{project_id}/fmea-templates"): "登记出表模板",
    ("GET", "/api/delivery/projects/{project_id}/fmea-templates/{template_id}"): "读一个出表模板",
    ("POST", "/api/delivery/projects/{project_id}/fmea-templates/{template_id}/{version}/approve"): "批准出表模板",
    ("POST", "/api/delivery/projects/{project_id}/documents/upload"): "往项目上传资料",
    ("GET", "/api/delivery/projects/{project_id}/tasks"): "列出后台任务",
    ("POST", "/api/delivery/projects/{project_id}/tasks"): "新建后台任务",
    ("GET", "/api/delivery/tasks/{task_id}"): "读一个任务",
    ("POST", "/api/delivery/tasks/{task_id}/heartbeat"): "任务心跳",
    ("POST", "/api/delivery/tasks/{task_id}/cancel"): "取消任务",
    ("POST", "/api/delivery/tasks/{task_id}/retry"): "重试任务",
    ("POST", "/api/delivery/tasks/batch-retry"): "批量重试任务",
    ("POST", "/api/delivery/tasks/{task_id}/assign"): "指派任务",
    ("GET", "/api/delivery/tasks/{task_id}/comments"): "读任务评论",
    ("POST", "/api/delivery/tasks/{task_id}/comments"): "写任务评论",
    ("GET", "/api/delivery/review-queue"): "待审队列",
    ("GET", "/api/delivery/source-assets"): "列出原始页图",
    ("GET", "/api/delivery/documents"): "列出资料版本",
    ("POST", "/api/delivery/documents/intake"): "解析进候选稿",
    ("POST", "/api/delivery/documents/intake/ocr-result"): "回写识图取字结果",
    ("GET", "/api/delivery/documents/{version_id}"): "读一个资料版本",
    ("GET", "/api/delivery/documents/{version_id}/review-package"): "整页审核包",
    ("POST", "/api/delivery/documents/{version_id}/review"): "审核资料",
    ("POST", "/api/delivery/documents/{version_id}/publish"): "发布正式资料",
    ("POST", "/api/delivery/documents/{version_id}/revise"): "改稿出新版本",
    ("POST", "/api/delivery/documents/batch-review"): "批量审核资料",
    ("GET", "/api/delivery/documents/compare/{left_id}/{right_id}"): "对照两个资料版本",
    ("POST", "/api/delivery/documents/{document_id}/rollback"): "资料退回旧版",
    ("GET", "/api/delivery/documents-search"): "在已发布资料里检索",
    ("GET", "/api/delivery/documents-index/status"): "检索索引状态",
    ("POST", "/api/delivery/documents-index/rebuild"): "重建检索索引",
    ("GET", "/api/delivery/documents/ocr-jobs/{job_id}"): "读识图任务",
    ("POST", "/api/delivery/documents/ocr-jobs/{job_id}/run"): "跑识图任务",
    ("POST", "/api/delivery/documents/ocr-jobs/{job_id}/retry-pages"): "重试失败页",
    ("GET", "/api/delivery/documents/source-assets/{asset_id}/pages/{page}"): "打开某一页原图",
    ("GET", "/api/delivery/evidence/{evidence_id}/open"): "按证据号回到原页",
    ("GET", "/api/delivery/graphs"): "列出图谱版本",
    ("POST", "/api/delivery/graphs/candidates"): "提交图谱候选，必须带证据",
    ("POST", "/api/delivery/graphs/extract"): "服务端抽图，需已登记提供方",
    ("GET", "/api/delivery/graphs/{graph_version_id}"): "读一个图谱版本",
    ("GET", "/api/delivery/graphs/{graph_version_id}/view"): "看图节点和边",
    ("GET", "/api/delivery/graphs/{graph_version_id}/statements"): "列出图上每条关系",
    ("POST", "/api/delivery/graphs/{graph_version_id}/statement-reviews"): "逐条审关系",
    ("POST", "/api/delivery/graphs/{graph_version_id}/review"): "审核整张图",
    ("POST", "/api/delivery/graphs/{graph_version_id}/publish"): "发布正式图",
    ("GET", "/api/delivery/graphs/{graph_version_id}/export"): "导出三元组",
    ("GET", "/api/delivery/graphs/{graph_version_id}/path"): "查两点之间的路径",
    ("GET", "/api/delivery/graphs/{graph_version_id}/evidence-audit"): "核对关系是否绑回原页",
    ("POST", "/api/delivery/graphs/{graph_version_id}/query"): "在已发布图上提问",
    ("POST", "/api/delivery/graphs/{graph_version_id}/compare-rag"): "同一题对照普通检索",
    ("POST", "/api/delivery/graphs/{graph_version_id}/resync"): "把已发布图同步进图库",
    ("GET", "/api/delivery/graphs/compare/{left_id}/{right_id}"): "对照两个图谱版本",
    ("POST", "/api/delivery/graphs/rollback"): "图谱退回旧版",
    ("GET", "/api/delivery/graphs-active/status"): "当前生效图状态",
    ("GET", "/api/delivery/graphs/{graph_version_id}/community-summaries"): "读社区摘要",
    ("POST", "/api/delivery/graphs/{graph_version_id}/community-summaries"): "写社区摘要",
    ("GET", "/api/delivery/fmea/templates"): "列出出表模板",
    ("GET", "/api/delivery/fmea/templates/{template_id}"): "读一个出表模板",
    ("GET", "/api/delivery/fmea/tasks"): "列出出表任务",
    ("POST", "/api/delivery/fmea/tasks"): "从已发布图生成表",
    ("GET", "/api/delivery/fmea/tasks/{task_id}"): "读一张表",
    ("POST", "/api/delivery/fmea/tasks/{task_id}/field-reviews"): "逐字段审核",
    ("POST", "/api/delivery/fmea/tasks/{task_id}/review"): "整表审核",
    ("POST", "/api/delivery/fmea/tasks/{task_id}/publish"): "发布正式表",
    ("GET", "/api/delivery/fmea/tasks/{task_id}/export"): "导出 JSON / CSV / DOCX",
    ("GET", "/api/delivery/fmea/tasks/{task_id}/export-verify"): "核对三份导出行数",
    ("GET", "/api/delivery/fmea/tasks/{task_id}/reviews"): "读审核记录",
    ("GET", "/api/delivery/fmea/tasks/{task_id}/status-history"): "读状态历史",
    ("POST", "/api/delivery/fmea/tasks/{task_id}/feedback"): "记下问题并指回上游",
    ("GET", "/api/delivery/fmea/tasks/{task_id}/feedback"): "列出反馈",
    ("POST", "/api/delivery/fmea/feedback/{feedback_id}/remediate"): "按反馈重跑上游",
    ("GET", "/api/delivery/fmea/feedback/{feedback_id}/runs"): "看整改跑次",
    ("GET", "/api/delivery/acceptance/status"): "验收门禁总状态",
    ("GET", "/api/delivery/acceptance/schema"): "验收材料格式",
    ("GET", "/api/delivery/acceptance/artifacts/{artifact_key}"): "读一份验收材料",
    ("PUT", "/api/delivery/acceptance/artifacts/{artifact_key}"): "写入验收材料",
    ("GET", "/api/delivery/acceptance/candidates/{gate_id}"): "读某门禁候选材料",
    ("POST", "/api/delivery/acceptance/evidence"): "上传验收截图或附件",
    ("GET", "/api/delivery/acceptance/evidence/{reference:path}"): "下载验收附件",
    ("POST", "/api/delivery/acceptance/build-package"): "生成验收包",
    ("GET", "/api/delivery/acceptance/handoff"): "读交接说明",
    ("GET", "/api/delivery/acceptance/package-files/{package_file}"): "下载验收包文件",
    ("GET", "/api/graphrag/communities"): "列出图社区",
    ("POST", "/api/graphrag/community/detect"): "做社区划分",
    ("POST", "/api/graphrag/community/summarize"): "给社区写摘要，需模型",
    ("POST", "/api/graphrag/search/global"): "按社区摘要做全局问答，需模型",
    ("POST", "/api/graphrag/import"): "把浏览器里的图写入图库",
    ("POST", "/api/graphrag/reset"): "清空运行图库",
    ("POST", "/api/graphrag/stats"): "图库统计",
    ("GET", "/api/graphrag/export"): "导出图库 JSON",
    ("GET", "/api/graphrag/triage"): "列出问答质检记录",
    ("GET", "/api/graphrag/triage/{triage_id}"): "读一条质检",
    ("POST", "/api/graphrag/triage/{triage_id}/review"): "审一条质检",
    ("POST", "/api/graphrag/triage/{triage_id}/promote"): "把质检升成回归题",
    ("GET", "/api/graphrag/triage/analytics"): "质检汇总",
    ("GET", "/api/graphrag/triage/export"): "导出质检记录",
    ("GET", "/api/memory/sessions"): "列出对话",
    ("POST", "/api/memory/sessions"): "建立或复用对话",
    ("GET", "/api/memory/sessions/{session_id}/messages"): "读对话消息",
    ("POST", "/api/memory/sessions/{session_id}/messages"): "追加一条消息",
    ("POST", "/api/memory/sessions/{session_id}/turns"): "追加一轮问答",
    ("GET", "/api/memory/sessions/{session_id}/context"): "取多轮上下文",
    ("DELETE", "/api/memory/sessions/{session_id}/messages"): "清空对话消息",
    ("DELETE", "/api/memory/sessions/{session_id}"): "删除对话",
    ("POST", "/api/v1/query"): "版本化问答，一次返回",
    ("POST", "/api/v1/query/stream"): "版本化问答，边生成边推",
    ("POST", "/api/v1/fmea/analyses/{analysis_id}/revisions"): "组装一版分析稿",
    ("GET", "/api/v1/fmea/revisions/{revision_id}"): "读分析稿",
    ("GET", "/api/v1/fmea/revisions/{revision_id}/readiness"): "看出稿是否齐",
    ("POST", "/api/v1/fmea/revisions/{revision_id}/readiness-suggestion-runs"): "给出稿缺口建议",
    ("POST", "/api/v1/fmea/revisions/{revision_id}/approval-submissions"): "提交审批",
    ("GET", "/api/v1/fmea/revisions/{revision_id}/approval-events"): "审批事件",
    ("POST", "/api/v1/fmea/approval-submissions/{submission_id}/approvals"): "批准分析稿",
    ("POST", "/api/v1/fmea/approval-submissions/{submission_id}/rejections"): "驳回分析稿",
    ("POST", "/api/v1/fmea/approvals/{approval_id}/withdrawals"): "撤回批准",
    ("POST", "/api/v1/fmea/revisions/{revision_id}/publications"): "发布分析稿",
    ("GET", "/api/v1/fmea/publications/{publication_id}"): "读发布体",
    ("GET", "/api/v1/fmea/publications/{publication_id}/snapshot"): "读发布快照",
    ("GET", "/api/v1/fmea/publications/{publication_id}/lifecycle-events"): "发布生命周期",
    ("POST", "/api/v1/fmea/publications/{publication_id}/withdrawals"): "撤回发布",
    ("POST", "/api/v1/fmea/publications/{publication_id}/supersessions"): "用新发布替换旧发布",
    ("POST", "/api/v1/fmea/template-drafts"): "新建模板草稿",
    ("POST", "/api/v1/fmea/template-drafts/{draft_id}/patch-runs"): "试改模板",
    ("GET", "/api/v1/fmea/template-patches/{patch_id}"): "读模板改动",
    ("POST", "/api/v1/fmea/template-patches/{patch_id}/acceptance"): "接受模板改动",
    ("POST", "/api/v1/fmea/template-patches/{patch_id}/rejection"): "拒绝模板改动",
    ("POST", "/api/v1/fmea/revisions/{revision_id}/migration-dry-runs"): "模板迁移试跑",
    ("POST", "/api/v1/fmea/migration-reports/{report_id}/confirmations"): "确认模板迁移",
    ("POST", "/api/v1/fmea/revisions/{revision_id}/export-runs"): "按修订导出",
    ("POST", "/api/v1/fmea/revisions/{revision_id}/export-narrative-runs"): "导出说明文字",
    ("GET", "/api/v1/fmea/export-runs/{run_id}"): "读导出任务",
    ("GET", "/api/v1/fmea/export-artifacts/{artifact_id}"): "取导出文件",
    ("GET", "/api/v1/fmea/rows/{row_id}/review-context"): "一行的审核上下文",
    ("POST", "/api/v1/fmea/rows/{row_id}/review-suggestion-runs"): "给一行出审核建议",
    ("GET", "/api/v1/fmea/review-suggestion-runs/{run_id}"): "读审核建议任务",
    ("GET", "/api/v1/fmea/rows/{row_id}/review-suggestions"): "列出一行的建议",
    ("POST", "/api/v1/fmea/rows/{row_id}/review-decisions"): "写下审核决定",
    ("GET", "/api/v1/fmea/rows/{row_id}/review-decisions"): "读审核决定",
    ("GET", "/api/v1/fmea/rows/{row_id}/risk"): "读一行风险",
    ("POST", "/api/v1/fmea/rows/{row_id}/risk-proposal-runs"): "提出风险分",
    ("GET", "/api/v1/fmea/risk-proposal-runs/{run_id}"): "读风险建议任务",
    ("POST", "/api/v1/fmea/rows/{row_id}/risk-confirmations"): "确认风险分",
    ("POST", "/api/v1/fmea/rows/{row_id}/risk-rejections"): "驳回风险分",
    ("POST", "/api/v1/fmea/analyses/{analysis_id}/propagation-runs"): "做故障传播分析",
    ("GET", "/api/v1/fmea/propagation-runs/{run_id}"): "读传播任务",
    ("GET", "/api/v1/fmea/propagation-graphs/{graph_revision_id}"): "读传播图",
    ("GET", "/api/v1/fmea/propagation-graphs/{graph_revision_id}/paths"): "列传播路径",
    ("POST", "/api/v1/fmea/propagation-graphs/{graph_revision_id}/reviews"): "审传播图",
    ("POST", "/api/v1/fmea/assistance/analysis-scope-runs"): "建议分析范围",
    ("GET", "/api/v1/fmea/assistance/suggestions/{suggestion_id}"): "读一条辅助建议",
    ("POST", "/api/v1/fmea/assistance/suggestions/{suggestion_id}/decisions"): "采纳或驳回辅助建议",
}

FAMILIES = [
    ("打开与文档", lambda m, p: p in {"/", "/api/health", "/docs", "/docs/oauth2-redirect", "/openapi.json", "/redoc"}),
    ("工作台入库检索", lambda m, p: p.startswith("/api/") and not p.startswith(("/api/delivery", "/api/graphrag", "/api/memory", "/api/v1", "/api/retrieval"))),
    ("检索策略与登录", lambda m, p: p.startswith("/api/retrieval/")),
    ("交付 · 项目与任务", lambda m, p: p.startswith("/api/delivery/projects") or p.startswith("/api/delivery/tasks") or p in {"/api/delivery/identity", "/api/delivery/review-queue", "/api/delivery/source-assets"}),
    ("交付 · 资料", lambda m, p: p.startswith("/api/delivery/documents") or p.startswith("/api/delivery/evidence") or p.startswith("/api/delivery/documents-")),
    ("交付 · 图谱", lambda m, p: p.startswith("/api/delivery/graphs")),
    ("交付 · 出表", lambda m, p: p.startswith("/api/delivery/fmea")),
    ("交付 · 验收", lambda m, p: p.startswith("/api/delivery/acceptance")),
    ("运行图与质检", lambda m, p: p.startswith("/api/graphrag")),
    ("对话记忆", lambda m, p: p.startswith("/api/memory")),
    ("版本化问答", lambda m, p: p.startswith("/api/v1/query")),
    ("出表治理 v1", lambda m, p: p.startswith("/api/v1/fmea")),
]


def load_routes() -> list[tuple[str, str, str]]:
    rows: list[tuple[str, str, str]] = []
    for line in ROUTE_FILE.read_text(encoding="utf-8").splitlines():
        if not line or line.startswith("count="):
            continue
        method, path, name = line.split("\t")
        rows.append((method, path, name))
    return rows


def describe(method: str, path: str, name: str) -> str:
    return DESC.get((method, path), name.replace("_", " "))


def family_of(method: str, path: str) -> str:
    for title, pred in FAMILIES:
        if pred(method, path):
            return title
    return "其他"


def _font(run, size, bold=False):
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    run.font.bold = bold
    run.font.name = FONT
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", FONT)


def text(slide, l, t, w, h, s, size=16, bold=False):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, part in enumerate(s.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = part
        _font(run, size, bold)
    return box


def footer(slide, s: str):
    text(slide, Inches(0.45), Emu(FOOT_TOP), Inches(12.4), Emu(620000), s, size=13)


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def style_cell(cell, value, size, bold=False):
    cell.fill.solid()
    cell.fill.fore_color.rgb = WHITE
    from pptx.oxml import parse_xml

    for border in ("lnL", "lnR", "lnT", "lnB"):
        ln = cell._tc.get_or_add_tcPr().find(qn(f"a:{border}"))
        if ln is None:
            cell._tc.get_or_add_tcPr().append(
                parse_xml(
                    f'<a:{border} w="6350" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
                    f'<a:solidFill><a:srgbClr val="000000"/></a:solidFill></a:{border}>'
                )
            )
    tf = cell.text_frame
    tf.word_wrap = True
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = value
    _font(run, size, bold)
    body = tf._txBody.find(qn("a:bodyPr"))
    if body is not None:
        body.set("anchor", "ctr")


def add_table(slide, rows, top, heights, widths, size=13):
    table = slide.shapes.add_table(
        len(rows), len(rows[0]), Inches(0.45), Inches(top), Inches(12.4), Inches(heights)
    ).table
    for i, w in enumerate(widths):
        table.columns[i].width = Inches(w)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            style_cell(table.cell(r, c), val, size if r else 12, bold=(r == 0 or c == 0))
    return table


def add_title_body(prs, title, body, foot=""):
    s = blank(prs)
    text(s, Inches(0.45), Inches(0.28), Inches(12.4), Inches(0.45), title, size=22)
    text(s, Inches(0.45), Inches(1.05), Inches(12.4), Inches(5.20), body, size=20)
    if foot:
        footer(s, foot)
    return s


def add_table_slide(prs, title, rows, widths, foot="", size=12, top=0.85, heights=5.55):
    s = blank(prs)
    text(s, Inches(0.45), Inches(0.22), Inches(12.4), Inches(0.40), title, size=20)
    add_table(s, rows, top, heights, widths, size=size)
    if foot:
        footer(s, foot)
    return s


def chunks(items, n):
    for i in range(0, len(items), n):
        yield items[i : i + n]


def build_ppt(routes: list[tuple[str, str, str]]) -> Presentation:
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    if prs.slides:
        r_id = prs.slides._sldIdLst[0].get(qn("r:id"))
        prs.part.drop_rel(r_id)
        prs.slides._sldIdLst.remove(prs.slides._sldIdLst[0])

    s = blank(prs)
    text(s, Inches(0.45), Inches(2.00), Inches(12.4), Inches(0.70), "PowerRAG 整合版", size=32, bold=True)
    text(
        s,
        Inches(0.45),
        Inches(2.80),
        Inches(12.4),
        Inches(1.80),
        "全项目技术链路与接口\n现网 218 条。浏览器打开，不走桌面壳。",
        size=22,
    )
    footer(s, "2026-09-23    现网：http://127.0.0.1:8000    接口文档：/docs")

    add_title_body(
        prs,
        "项目是什么",
        "本地工作台。手册和扫描件进来，读成带页码的证据，再检索、构图、问答、出故障分析表。\n"
        "打开方式：双击 PowerRAG.exe，或 python api_server/current_console/server.py。\n"
        "地址 http://127.0.0.1:8000。关掉状态窗口即停服务。\n"
        "仓库里的 Electron 目录只保留，现网不用。RAG交付 是历史快照，接口是子集。",
        "现网入口只有 current_console。版本号见 /api/health。",
    )

    add_table_slide(
        prs,
        "仓库怎么分层",
        [
            ("层", "目录", "管什么"),
            ("界面", "frontend_app/current_console", "五页工作台加设置抽屉"),
            ("接口", "api_server/current_console", "现网 218 条，版本 2.1.0"),
            ("解析", "data_pipeline", "读文件、识图取字、保住页码"),
            ("入库", "knowledge_base / storage_layer", "正式版本、向量库、治理库"),
            ("图谱", "kg_pipeline / storage_layer", "抽关系、社区、图库"),
            ("问答", "rag_orchestrator", "普通检索、沿图回答、对照"),
            ("出表", "rag_orchestrator / fmea_*", "从已发布图出可审核表"),
            ("验收", "evaluation", "门禁材料和验收包"),
            ("启动", "desktop_launcher", "Windows 可执行文件，仍开系统浏览器"),
        ],
        [1.6, 4.4, 6.4],
        "治理库记审核和版本。向量库做检索。图库做沿边问答。三套不混写。",
        size=13,
        heights=5.50,
    )

    add_title_body(
        prs,
        "两套编号",
        "界面从左到右：M1 资料接入 → M2 知识组织 → M3 图谱问答 → M4 可信交付 → M5 验收。\n"
        "交付治理另一套：解析 → 入库 → 构图 → 出表。后一段只吃前一段已经发布的结果。\n"
        "界面 M4 可信交付，对应治理这条解析到出表。界面 M5 是验收台，不是出表本身。\n"
        "设置和运行总览在右上角抽屉，不占主流水线。",
        "FMEA：故障模式与影响分析。出表字段必须能回到原文页码。",
    )

    add_title_body(
        prs,
        "全项目一条线",
        "资料进来，读成带页码的证据。\n"
        "审过的稿做成可检索正式版本，能对照、能退回。\n"
        "只从已发布资料构图，关系必须绑回原页。这是相对把资料直接丢给大模型作答的优势。\n"
        "问答三条路：关键词检索、沿已发布图回答、同一题对照。原文没有的要能拒绝。\n"
        "只从已发布图出故障分析表，每个字段必须带证据。再走验收门禁和整包进出。",
        "解析 — 入库 — 图谱 — 问答 — 出表 — 验收",
    )

    add_table_slide(
        prs,
        "界面五页",
        [
            ("页", "做什么", "主要接口族"),
            ("M1 资料接入", "上传、解析、写入检索库", "/api/upload /process /ingest /search"),
            ("M2 知识组织", "构图、社区、导入导出图", "/api/graphrag/*"),
            ("M3 图谱问答", "提问、多轮、对照", "/api/query  /api/v1/query  /api/memory"),
            ("M4 可信交付", "项目、审核、发布、出表", "/api/delivery/*"),
            ("M5 验收", "门禁材料、证据、验收包", "/api/delivery/acceptance/*"),
            ("抽屉 · 总览", "运行状态", "/api/health  /api/stats"),
            ("抽屉 · 设置", "集合、策略、身份、日志", "/api/retrieval/policies  /api/logs"),
        ],
        [2.4, 4.4, 5.6],
        "page-search、page-kg 是旧页，不挂进现网导航。",
        size=14,
        heights=5.40,
    )

    grouped: dict[str, list[tuple[str, str, str]]] = {title: [] for title, _ in FAMILIES}
    grouped["其他"] = []
    for method, path, name in routes:
        grouped[family_of(method, path)].append((method, path, describe(method, path, name)))

    summary = [("接口族", "前缀", "条数")]
    prefix_hint = {
        "打开与文档": "/  /docs  /openapi.json",
        "工作台入库检索": "/api/upload /process /ingest /query …",
        "检索策略与登录": "/api/retrieval/policies",
        "交付 · 项目与任务": "/api/delivery/projects  /tasks",
        "交付 · 资料": "/api/delivery/documents",
        "交付 · 图谱": "/api/delivery/graphs",
        "交付 · 出表": "/api/delivery/fmea",
        "交付 · 验收": "/api/delivery/acceptance",
        "运行图与质检": "/api/graphrag",
        "对话记忆": "/api/memory",
        "版本化问答": "/api/v1/query",
        "出表治理 v1": "/api/v1/fmea",
        "其他": "—",
    }
    for title, _ in FAMILIES:
        summary.append((title, prefix_hint[title], str(len(grouped[title]))))
    if grouped["其他"]:
        summary.append(("其他", "—", str(len(grouped["其他"]))))
    summary.append(("合计", "现网 create_app 实扫", str(len(routes))))
    add_table_slide(
        prs,
        "现网接口族",
        summary,
        [3.2, 6.6, 2.6],
        "权威清单以运行中的 /openapi.json 为准。本表来自 2026-09-23 实扫，218 条。",
        size=14,
        heights=5.50,
    )

    for title, _ in FAMILIES:
        items = grouped[title]
        for i, batch in enumerate(chunks(items, 10), start=1):
            rows = [("方法", "路径", "做什么")]
            rows.extend(batch)
            suffix = f"（{i}）" if len(items) > 10 else ""
            add_table_slide(
                prs,
                f"{title}{suffix}  ·  {len(items)} 条",
                rows,
                [1.4, 6.5, 4.5],
                "路径里的花括号是变量。POST 写，GET 读，PUT / PATCH 改，DELETE 删。",
                size=11,
                top=0.78,
                heights=5.60,
            )

    add_title_body(
        prs,
        "受控失败，不是要绕开的错",
        "扫描件缺字、缺页码、缺审核，不能发布正式资料。\n"
        "未发布资料不能构图。类型或关系不在规则里、证据不在所选版本，不能发布图。\n"
        "未发布图不能出表。字段对不上或互相打架，保持空，不编严重度、频度、探测度。\n"
        "未发布的表不能当正式导出。记下反馈和真的重跑上游，不是同一件事。",
        "S / O / D：严重度、频度、探测度。没有批准的评分政策就不写这三项。",
    )
    return prs


def write_md(routes: list[tuple[str, str, str]]) -> str:
    grouped: dict[str, list[tuple[str, str, str]]] = {title: [] for title, _ in FAMILIES}
    grouped["其他"] = []
    for method, path, name in routes:
        grouped[family_of(method, path)].append((method, path, describe(method, path, name)))

    lines = [
        "# PowerRAG 整合版 · 全项目技术链路与接口",
        "",
        "日期：2026-09-23  ",
        "现网入口：`http://127.0.0.1:8000`  ",
        "接口文档：`http://127.0.0.1:8000/docs`  ",
        "权威清单：`http://127.0.0.1:8000/openapi.json`  ",
        f"实扫条数：**{len(routes)}**（`api_server/current_console` `create_app`）",
        "",
        "## 项目是什么",
        "",
        "本地工作台。手册和扫描件进来，读成带页码的证据，再检索、构图、问答、出故障分析表。",
        "打开：双击 `desktop_launcher/dist/PowerRAG/PowerRAG.exe`，或 `python api_server/current_console/server.py`。",
        "仓库里的 Electron 目录只保留，现网不用。`RAG交付/` 是历史快照，接口是子集。",
        "",
        "## 仓库分层",
        "",
        "| 层 | 目录 | 管什么 |",
        "|---|---|---|",
        "| 界面 | `frontend_app/current_console` | 五页工作台加设置抽屉 |",
        "| 接口 | `api_server/current_console` | 现网 218 条，版本 2.1.0 |",
        "| 解析 | `data_pipeline` | 读文件、识图取字、保住页码 |",
        "| 入库 | `knowledge_base` / `storage_layer` | 正式版本、向量库、治理库 |",
        "| 图谱 | `kg_pipeline` / `storage_layer` | 抽关系、社区、图库 |",
        "| 问答 | `rag_orchestrator` | 普通检索、沿图回答、对照 |",
        "| 出表 | `rag_orchestrator` / `fmea_*` | 从已发布图出可审核表 |",
        "| 验收 | `evaluation` | 门禁材料和验收包 |",
        "| 启动 | `desktop_launcher` | Windows 可执行文件，仍开系统浏览器 |",
        "",
        "治理库记审核和版本。向量库做检索。图库做沿边问答。三套不混写。",
        "",
        "## 两套编号",
        "",
        "- 界面：M1 资料接入 → M2 知识组织 → M3 图谱问答 → M4 可信交付 → M5 验收。",
        "- 交付治理：解析 → 入库 → 构图 → 出表。后一段只吃前一段已经发布的结果。",
        "- 界面 M4 对应治理这条解析到出表。界面 M5 是验收台，不是出表本身。",
        "",
        "## 全项目一条线",
        "",
        "1. 资料进来，读成带页码的证据。",
        "2. 审过的稿做成可检索正式版本，能对照、能退回。",
        "3. 只从已发布资料构图，关系必须绑回原页。这是相对把资料直接丢给大模型作答的优势。",
        "4. 问答三条路：关键词检索、沿已发布图回答、同一题对照。原文没有的要能拒绝。",
        "5. 只从已发布图出故障分析表，每个字段必须带证据。",
        "6. 验收门禁和整包进出。",
        "",
        "## 界面五页",
        "",
        "| 页 | 做什么 | 主要接口族 |",
        "|---|---|---|",
        "| M1 资料接入 | 上传、解析、写入检索库 | `/api/upload` `/api/process` `/api/ingest` `/api/search` |",
        "| M2 知识组织 | 构图、社区、导入导出图 | `/api/graphrag/*` |",
        "| M3 图谱问答 | 提问、多轮、对照 | `/api/query` `/api/v1/query` `/api/memory` |",
        "| M4 可信交付 | 项目、审核、发布、出表 | `/api/delivery/*` |",
        "| M5 验收 | 门禁材料、证据、验收包 | `/api/delivery/acceptance/*` |",
        "| 抽屉 · 总览 | 运行状态 | `/api/health` `/api/stats` |",
        "| 抽屉 · 设置 | 集合、策略、身份、日志 | `/api/retrieval/policies` `/api/logs` |",
        "",
        "## 现网全部接口",
        "",
    ]
    for title, _ in FAMILIES:
        items = grouped[title]
        lines.append(f"### {title}（{len(items)}）")
        lines.append("")
        lines.append("| 方法 | 路径 | 做什么 |")
        lines.append("|---|---|---|")
        for method, path, desc in items:
            lines.append(f"| `{method}` | `{path}` | {desc} |")
        lines.append("")
    if grouped["其他"]:
        lines.append("### 其他")
        lines.append("")
        lines.append("| 方法 | 路径 | 做什么 |")
        lines.append("|---|---|---|")
        for method, path, desc in grouped["其他"]:
            lines.append(f"| `{method}` | `{path}` | {desc} |")
        lines.append("")
    lines.extend(
        [
            "## 受控失败",
            "",
            "扫描件缺字、缺页码、缺审核，不能发布正式资料。未发布资料不能构图。",
            "类型或关系不在规则里、证据不在所选版本，不能发布图。未发布图不能出表。",
            "字段对不上或互相打架，保持空，不编严重度、频度、探测度。未发布的表不能当正式导出。",
            "",
        ]
    )
    return "\n".join(lines)


def write_html(md_text: str, routes: list[tuple[str, str, str]]) -> str:
    grouped: dict[str, list[tuple[str, str, str]]] = {title: [] for title, _ in FAMILIES}
    grouped["其他"] = []
    for method, path, name in routes:
        grouped[family_of(method, path)].append((method, path, describe(method, path, name)))

    sections = []
    toc = []
    for idx, (title, _) in enumerate(FAMILIES, start=1):
        items = grouped[title]
        hid = f"f{idx}"
        toc.append(f'<a href="#{hid}">{escape(title)}（{len(items)}）</a>')
        rows = "".join(
            f"<tr><td>{escape(m)}</td><td><code>{escape(p)}</code></td><td>{escape(d)}</td></tr>"
            for m, p, d in items
        )
        sections.append(
            f'<h2 id="{hid}">{escape(title)} <small>{len(items)}</small></h2>'
            f"<table><thead><tr><th>方法</th><th>路径</th><th>做什么</th></tr></thead><tbody>{rows}</tbody></table>"
        )
    return f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8" />
  <title>PowerRAG 整合版 · 全项目技术链路与接口</title>
  <style>
    body {{ margin: 0; background: #fff; color: #000; font: 16px/1.6 "Microsoft YaHei", "PingFang SC", sans-serif; }}
    main {{ max-width: 1100px; margin: 0 auto; padding: 32px 28px 80px; }}
    h1 {{ font-size: 28px; font-weight: 600; margin: 0 0 8px; }}
    h2 {{ font-size: 20px; margin: 36px 0 12px; }}
    h3 {{ font-size: 18px; margin: 28px 0 10px; }}
    p, li {{ font-size: 16px; }}
    .lead {{ font-size: 18px; }}
    .meta {{ color: #333; margin-bottom: 24px; }}
    .toc {{ display: flex; flex-wrap: wrap; gap: 8px 16px; margin: 16px 0 28px; }}
    .toc a {{ color: #000; }}
    table {{ width: 100%; border-collapse: collapse; margin: 8px 0 24px; }}
    th, td {{ border: 1px solid #000; padding: 7px 8px; text-align: left; vertical-align: top; }}
    th {{ font-weight: 600; }}
    code {{ font-family: Consolas, "Courier New", monospace; font-size: 13px; }}
    small {{ font-weight: 400; }}
  </style>
</head>
<body>
<main>
  <h1>PowerRAG 整合版 · 全项目技术链路与接口</h1>
  <p class="meta">2026-09-23 · 现网 {len(routes)} 条 · http://127.0.0.1:8000 · /docs</p>
  <p class="lead">本地工作台。手册和扫描件进来，读成带页码的证据，再检索、构图、问答、出故障分析表。现网入口是 current_console。Electron 不用。RAG交付 是历史快照。</p>
  <h2>目录</h2>
  <div class="toc">{"".join(toc)}</div>
  <h2>一条线</h2>
  <p>资料进来，读成带页码的证据。审过再入库。只从已发布资料构图，关系绑回原页——相对把资料直接丢给大模型作答的优势。问答走检索或沿图。只从已发布图出表，字段带证据。最后走验收。</p>
  <h2>两套编号</h2>
  <p>界面：M1 资料接入 → M2 知识组织 → M3 图谱问答 → M4 可信交付 → M5 验收。<br>治理：解析 → 入库 → 构图 → 出表。界面 M4 对应治理全段。界面 M5 是验收台。</p>
  {"".join(sections)}
</main>
</body>
</html>
"""


def _save(prs: Presentation, dest: Path) -> Path:
    dest.parent.mkdir(parents=True, exist_ok=True)
    tmp = dest.with_name("_tmp_" + dest.name)
    prs.save(str(tmp))
    try:
        dest.write_bytes(tmp.read_bytes())
        tmp.unlink(missing_ok=True)
        return dest
    except OSError:
        print("locked", dest, "->", tmp)
        return tmp


def main() -> None:
    routes = load_routes()
    if len(routes) != 218:
        raise SystemExit(f"expected 218 routes, got {len(routes)}")
    missing = [(m, p) for m, p, _ in routes if (m, p) not in DESC]
    if missing:
        raise SystemExit("missing DESC: " + "; ".join(f"{m} {p}" for m, p in missing))
    prs = build_ppt(routes)
    md = write_md(routes)
    html = write_html(md, routes)
    for dest in (MD, USER_MD):
        dest.write_text(md, encoding="utf-8")
    for dest in (HTML, USER_HTML):
        dest.write_text(html, encoding="utf-8")
    saved = [_save(prs, OUT), _save(prs, USER_PPT)]
    print("routes", len(routes))
    print("slides", len(prs.slides))
    for path in saved + [HTML, USER_HTML, MD, USER_MD]:
        print(path)


if __name__ == "__main__":
    main()
