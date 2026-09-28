# -*- coding: utf-8 -*-
"""白底黑字。对照只做一张表，不拆页注水。"""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
OUT = HERE / "PowerRAG_8.8对照收尾.pptx"
USER_OUT = Path(r"D:\虚拟C盘\PowerRAG_8.8对照收尾.pptx")
HTML = HERE / "PowerRAG_8.8对照收尾.html"

BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
W = Inches(13.333)
H = Inches(7.5)
FONT = "微软雅黑"

ROWS = [
    ("", "8.15", "现在"),
    ("解析", "只有接口，合成样例能跑", "真实资料上传后，原页和解析在同一屏"),
    ("入库", "资料能进库、能搜", "改完出新稿，能对照旧稿，也能退回旧稿"),
    ("图谱", "抽图不够，还要回到原来的 RAG", "Schema 出图后，由 GraphRAG 沿路径回答"),
    ("FMEA", "字段带着来源", "故障分析表，出处写到页，导出 JSON / CSV / DOCX"),
]

INTRO_LEAD = "面向燃气轮机与动力装备资料的本地 GraphRAG 工作台。"
INTRO_BODY = (
    "产品就是这一个工作台，浏览器打开。\n"
    "手册、扫描件上传后，解析、入库、图谱问答都能做完。\n"
    "按燃气轮机与动力装备 Schema 抽出设备、部件、故障、原因、措施，关系落在证据边上。\n"
    "GraphRAG 沿路径问答。\n"
    "FMEA 是故障模式与影响分析：设备、故障、原因、措施做成一张表，出处写到页，导出 JSON / CSV / DOCX。"
)

DELIVER_BODY = (
    "一个本地工作台，浏览器打开就能用。\n"
    "用户上传资料，解析、入库、图谱问答走完。\n"
    "FMEA 故障分析表能出 JSON / CSV / DOCX。"
)
DELIVER_FOOT = "解析  —  入库  —  图谱问答"


def _font(run, size, bold=False):
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    run.font.bold = bold
    run.font.name = FONT
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", FONT)


def text(slide, l, t, w, h, s, size=16, bold=False):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, part in enumerate(s.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        run = p.add_run()
        run.text = part
        _font(run, size, bold)
    return box


def notes(slide, body: str):
    frame = slide.notes_slide.notes_text_frame
    if frame is None:
        return
    frame.text = body


def _drop_default(prs):
    if not prs.slides:
        return
    r_id = prs.slides._sldIdLst[0].get(qn("r:id"))
    prs.part.drop_rel(r_id)
    prs.slides._sldIdLst.remove(prs.slides._sldIdLst[0])


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def style_cell(cell, value, size, bold=False):
    cell.fill.solid()
    cell.fill.fore_color.rgb = WHITE
    for border in ("lnL", "lnR", "lnT", "lnB"):
        ln = cell._tc.get_or_add_tcPr().find(qn(f"a:{border}"))
        if ln is None:
            from pptx.oxml import parse_xml

            cell._tc.get_or_add_tcPr().append(
                parse_xml(
                    f'<a:{border} w="6350" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
                    f'<a:solidFill><a:srgbClr val="000000"/></a:solidFill></a:{border}>'
                )
            )
    tf = cell.text_frame
    tf.word_wrap = True
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = value
    _font(run, size, bold)
    body = tf._txBody.find(qn("a:bodyPr"))
    if body is not None:
        body.set("anchor", "ctr")


def add_table_slide(prs):
    s = blank(prs)
    table = s.shapes.add_table(len(ROWS), 3, Inches(0.45), Inches(1.15), Inches(12.4), Inches(5.20)).table
    table.columns[0].width = Inches(1.70)
    table.columns[1].width = Inches(5.35)
    table.columns[2].width = Inches(5.35)
    for r, row in enumerate(ROWS):
        for c, val in enumerate(row):
            style_cell(table.cell(r, c), val, 16 if r else 15, bold=(r == 0 or c == 0))


def add_intro(prs):
    s = blank(prs)
    text(s, Inches(0.45), Inches(1.15), Inches(12.4), Inches(0.80), INTRO_LEAD, size=22)
    text(s, Inches(0.45), Inches(2.15), Inches(12.4), Inches(4.40), INTRO_BODY, size=20)


def add_deliver(prs):
    s = blank(prs)
    text(s, Inches(0.45), Inches(0.28), Inches(12.4), Inches(0.40), "交什么", size=22)
    text(s, Inches(0.45), Inches(1.35), Inches(12.4), Inches(4.40), DELIVER_BODY, size=22)
    text(s, Inches(0.45), Inches(6.55), Inches(12.4), Inches(0.50), DELIVER_FOOT, size=16)


def add_walk(prs):
    s = blank(prs)
    text(s, Inches(0.45), Inches(0.28), Inches(12.4), Inches(0.40), "流程", size=22)
    text(s, Inches(0.45), Inches(1.50), Inches(12.4), Inches(4.40), CAP_BODY, size=22)


def write_html():
    body_rows = []
    for i, row in enumerate(ROWS):
        tag = "th" if i == 0 else "td"
        body_rows.append("<tr>" + "".join(f"<{tag}>{c}</{tag}>" for c in row) + "</tr>")
    table = "<table>" + "".join(body_rows) + "</table>"
    HTML.write_text(
        f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8" />
  <title>PowerRAG 对照 8.8 收尾</title>
  <style>
    html, body {{ margin: 0; height: 100%; background: #ececec; color: #000;
      font-family: "Microsoft YaHei", "PingFang SC", sans-serif; }}
    .stage {{ height: 100%; display: flex; align-items: center; justify-content: center; }}
    .deck {{ width: min(100vw, calc(100vh * 16 / 9)); height: min(100vh, calc(100vw * 9 / 16));
      background: #fff; position: relative; }}
    .slide {{ display: none; position: absolute; inset: 0; padding: 4.4% 4.2% 7%; box-sizing: border-box; }}
    .slide.on {{ display: block; }}
    h1 {{ font-size: 28px; font-weight: 400; margin: 0 0 3.6%; }}
    .lead {{ font-size: 26px; line-height: 1.55; margin: 0 0 0.8em; }}
    .copy {{ font-size: 22px; line-height: 1.7; margin: 0; }}
    table {{ width: 100%; border-collapse: collapse; font-size: 20px; }}
    th, td {{ border: 1px solid #000; padding: 14px 12px; text-align: left; font-weight: 400; }}
    th {{ font-weight: 600; }}
    td:first-child, th:first-child {{ width: 14%; font-weight: 600; }}
    .foot {{ position: absolute; left: 4.2%; right: 4.2%; bottom: 6.2%; font-size: 20px; }}
    .bar {{ position: fixed; right: 16px; bottom: 12px; font-size: 12px; color: #666; }}
  </style>
</head>
<body>
  <div class="stage"><div class="deck" id="deck"></div></div>
  <div class="bar" id="bar"></div>
  <script>
    const slides = [
      {{ t: "", html: `{table}`, f: "" }},
      {{ t: "",
        html: `<p class="lead">{INTRO_LEAD}</p>
          <p class="copy">{INTRO_BODY.replace(chr(10), "<br>")}</p>`,
        f: "" }},
      {{ t: "交什么",
        html: `<p class="copy">{DELIVER_BODY.replace(chr(10), "<br>")}</p>`,
        f: "{DELIVER_FOOT}" }}
    ];
    const deck = document.getElementById("deck");
    slides.forEach((s) => {{
      const el = document.createElement("section");
      el.className = "slide";
      el.innerHTML = (s.t ? `<h1>${{s.t}}</h1>` : "") + s.html + (s.f ? `<div class="foot">${{s.f}}</div>` : "");
      deck.appendChild(el);
    }});
    let i = 0;
    const pages = [...deck.children];
    const bar = document.getElementById("bar");
    const show = (n) => {{
      i = (n + pages.length) % pages.length;
      pages.forEach((p, k) => p.classList.toggle("on", k === i));
      bar.textContent = `${{String(i + 1).padStart(2, "0")}} / ${{String(pages.length).padStart(2, "0")}}  ·  ← →`;
    }};
    show(0);
    window.addEventListener("keydown", (e) => {{
      if (["ArrowRight", " ", "PageDown"].includes(e.key)) show(i + 1);
      if (["ArrowLeft", "PageUp"].includes(e.key)) show(i - 1);
    }});
  </script>
</body>
</html>
""",
        encoding="utf-8",
    )


def _close_if_open():
    try:
        import win32com.client

        app = win32com.client.Dispatch("PowerPoint.Application")
        targets = {OUT.resolve(), USER_OUT.resolve()}
        for i in range(app.Presentations.Count, 0, -1):
            pres = app.Presentations.Item(i)
            try:
                full = Path(pres.FullName).resolve()
            except Exception:
                continue
            if full in targets:
                pres.Save()
                pres.Close()
    except Exception:
        pass


def _save(prs, dest: Path) -> Path:
    tmp = dest.with_name("_tmp_" + dest.name)
    prs.save(tmp)
    try:
        dest.write_bytes(tmp.read_bytes())
        tmp.unlink(missing_ok=True)
        return dest
    except OSError:
        print("locked", dest, "->", tmp)
        return tmp


def build():
    _close_if_open()
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    _drop_default(prs)
    add_table_slide(prs)
    add_intro(prs)
    add_deliver(prs)
    write_html()
    _save(prs, OUT)
    return _save(prs, USER_OUT)


if __name__ == "__main__":
    path = build()
    print(path)
    print("slides", len(Presentation(str(path)).slides))
