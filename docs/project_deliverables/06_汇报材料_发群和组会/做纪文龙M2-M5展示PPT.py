# -*- coding: utf-8 -*-
"""白底黑字。纪文龙 M2–M5 上台稿，不把总验收写成通过。"""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
OUT = HERE / "纪文龙_M2-M5工作展示_20260923.pptx"

BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "微软雅黑"

WORK = [
    ("", "1.0 要求", "现在能交出去的"),
    ("M2 解析", "抽字/OCR，保住页码结构，对照原页审", "五类抽检过；清晰页关键字段核对过"),
    ("M3 入库", "审过再发布；检索、引用、回滚", "Recall@5=1.0；四类边界过；引用可回原块"),
    ("M4 图谱", "只从已发布资料抽关系，绑证据", "120 条关系审核过；18 条发布图绑证据；同题 4 问"),
    ("M5 FMEA", "已发布图+模板，字段带证据，人审导出", "3 条专家审过；JSON/CSV/DOCX 行数一致；不编 RPN"),
    ("界面闭环", "人在工作台上走完", "项目 jl-e2e 已建；资料/图/FMEA 已发布；回滚和恢复已跑通"),
    ("整包验收", "12 个门禁全过才算通过", "8 过，2 不过，2 未评"),
]

GATES = [
    ("门禁", "状态", "一句话"),
    ("AC-M2-01 / 02", "过", "页覆盖、五类抽检、清晰页核对"),
    ("AC-M3-01 / 02", "过", "Recall@5=1.0，四类边界过"),
    ("AC-M4-01 / 02", "过", "关系审核过，发布图绑证据"),
    ("AC-M5-01 / 02", "过", "字段有证据，三份导出一致"),
    ("AC-E2E-01 / 02", "不过", "差领域审核人签字。必须两个角色只点界面"),
    ("TECH-15 / 16", "未评", "测试负责人：参考机耗时、60MB 实传"),
]


def _font(run, size, bold=False):
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    run.font.bold = bold
    run.font.name = FONT
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", FONT)


def text(slide, l, t, w, h, s, size=16, bold=False):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, part in enumerate(s.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = part
        _font(run, size, bold)
    return box


def notes(slide, body: str):
    frame = slide.notes_slide.notes_text_frame
    if frame is None:
        return
    frame.text = body


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def style_cell(cell, value, size, bold=False):
    cell.fill.solid()
    cell.fill.fore_color.rgb = WHITE
    for border in ("lnL", "lnR", "lnT", "lnB"):
        ln = cell._tc.get_or_add_tcPr().find(qn(f"a:{border}"))
        if ln is None:
            from pptx.oxml import parse_xml

            cell._tc.get_or_add_tcPr().append(
                parse_xml(
                    f'<a:{border} w="6350" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
                    f'<a:solidFill><a:srgbClr val="000000"/></a:solidFill></a:{border}>'
                )
            )
    tf = cell.text_frame
    tf.word_wrap = True
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = value
    _font(run, size, bold)
    body = tf._txBody.find(qn("a:bodyPr"))
    if body is not None:
        body.set("anchor", "ctr")


def add_table(slide, rows, top, heights, widths, size=15):
    table = slide.shapes.add_table(
        len(rows),
        len(rows[0]),
        Inches(0.45),
        Inches(top),
        Inches(12.4),
        Inches(heights),
    ).table
    for i, w in enumerate(widths):
        table.columns[i].width = Inches(w)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            style_cell(table.cell(r, c), val, size if r else 14, bold=(r == 0 or c == 0))
    return table


def add_cover(prs):
    s = blank(prs)
    text(s, Inches(0.45), Inches(2.05), Inches(12.4), Inches(0.70), "M2–M5 工作展示", size=32, bold=True)
    text(
        s,
        Inches(0.45),
        Inches(2.90),
        Inches(12.4),
        Inches(1.40),
        "PowerRAG · 解析到可审核 FMEA\nM2 解析  ·  M3 入库  ·  M4 图谱  ·  M5 输出",
        size=22,
    )
    text(s, Inches(0.45), Inches(6.55), Inches(12.4), Inches(0.50), "2026-09-23", size=16)
    notes(s, "开场只报模块范围。")


def add_scope(prs):
    s = blank(prs)
    text(s, Inches(0.45), Inches(0.28), Inches(12.4), Inches(0.40), "1.0 分工", size=22)
    text(
        s,
        Inches(0.45),
        Inches(1.20),
        Inches(12.4),
        Inches(4.40),
        "M2、M3、M4、M5：解析到可审核 FMEA。\n"
        "M2 共同项：对照原页复核。\n"
        "M3 共同项、M6：编排和验收报告。\n"
        "TECH-15 / TECH-16：测试负责人项。",
        size=22,
    )
    text(s, Inches(0.45), Inches(6.55), Inches(12.4), Inches(0.50), "后一段只吃前一段已经发布的东西。", size=16)
    notes(s, "M1、M6、性能测试不并进本页范围。")


def add_line(prs):
    s = blank(prs)
    text(s, Inches(0.45), Inches(0.28), Inches(12.4), Inches(0.40), "一条线", size=22)
    text(
        s,
        Inches(0.45),
        Inches(1.35),
        Inches(12.4),
        Inches(4.40),
        "资料进来，读成带页码的证据。\n"
        "人审之后才发布正式版本。\n"
        "只从已发布资料构图，关系必须绑证据。\n"
        "只从已发布图出 FMEA，字段必须带证据。\n"
        "没有批准的评分政策，不编 S/O/D，不编 RPN。",
        size=22,
    )
    text(s, Inches(0.45), Inches(6.55), Inches(12.4), Inches(0.50), "解析  —  入库  —  图谱  —  FMEA", size=18)
    notes(s, "先讲约束，再点界面。图不够回退普通 RAG，原文没有的要能拒绝。")


def add_work(prs):
    s = blank(prs)
    text(s, Inches(0.45), Inches(0.28), Inches(12.4), Inches(0.40), "做成了什么", size=22)
    add_table(s, WORK, 0.85, 5.45, [1.70, 5.35, 5.35], size=14)
    text(s, Inches(0.45), Inches(6.55), Inches(12.4), Inches(0.50), "工程和金标门禁：8 项通过。", size=16)
    notes(s, "只读右边。FMEA 三份导出行数一致。")


def add_gates(prs):
    s = blank(prs)
    text(s, Inches(0.45), Inches(0.28), Inches(12.4), Inches(0.40), "门禁", size=22)
    add_table(s, GATES, 0.90, 5.20, [2.60, 1.20, 8.60], size=16)
    text(s, Inches(0.45), Inches(6.55), Inches(12.4), Inches(0.50), "12 个全过才算通过。现在 8 / 2 / 2。", size=16)
    notes(
        s,
        "E2E 操作已经齐，差领域审核人在界面签字。\n"
        "同一账号签两个角色，验收器不认。TECH 两项不并进本范围。",
    )


def add_demo(prs):
    s = blank(prs)
    text(s, Inches(0.45), Inches(0.28), Inches(12.4), Inches(0.40), "现在打开就能点", size=22)
    text(
        s,
        Inches(0.45),
        Inches(1.20),
        Inches(12.4),
        Inches(4.60),
        "浏览器  http://127.0.0.1:8000\n"
        "进 M4 可信交付，项目选 jl-e2e-20260923。\n"
        "资料 lube-oil-filter 已发布到 v5。\n"
        "图谱 graph:v1，7 条语句，已发布。\n"
        "FMEA 已发布，可导出 CSV / 正式 DOCX。\n"
        "恢复项目 jl-e2e-restored-20260923 已经在。",
        size=22,
    )
    text(s, Inches(0.45), Inches(6.55), Inches(12.4), Inches(0.50), "系统浏览器打开。", size=16)
    notes(
        s,
        "现场只点这一条：资料台账 → 图 → FMEA 导出。\n"
        "回滚看「回滚到 v4」。验收台切 M5。",
    )


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    add_cover(prs)
    add_scope(prs)
    add_line(prs)
    add_work(prs)
    add_gates(prs)
    add_demo(prs)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
