from __future__ import annotations

from pathlib import Path
from math import atan2, cos, sin, pi

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "diagrams" / "key_stage_flows"
DOCX_OUT = ROOT / "docs" / "PowerRAG_关键环节流程图合集_2026-07-06.docx"
MD_OUT = ROOT / "docs" / "PowerRAG_关键环节流程图合集_2026-07-06.md"


def load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        Path("C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/simsun.ttc"),
    ]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


FONT_TITLE = load_font(50, True)
FONT_SUBTITLE = load_font(24)
FONT_SECTION = load_font(28, True)
FONT_BOX_TITLE = load_font(24, True)
FONT_BOX = load_font(19)
FONT_SMALL = load_font(17)
FONT_TAG = load_font(18, True)


def text_width(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.ImageFont) -> int:
    if not text:
        return 0
    box = draw.textbbox((0, 0), text, font=font)
    return box[2] - box[0]


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.ImageFont, max_width: int) -> list[str]:
    lines: list[str] = []
    current = ""
    break_chars = " ，、；;，。,.：:()/（）"
    for char in text:
        trial = current + char
        if text_width(draw, trial, font) <= max_width or not current:
            current = trial
        else:
            break_at = max(current.rfind(c) for c in break_chars)
            if break_at > 0:
                line = current[: break_at + 1].rstrip()
                rest = current[break_at + 1 :].lstrip()
                if line:
                    lines.append(line)
                    current = rest + char
                else:
                    lines.append(current)
                    current = char
            else:
                lines.append(current)
                current = char
    if current:
        lines.append(current)
    return lines


def rounded(
    draw: ImageDraw.ImageDraw,
    rect: tuple[int, int, int, int],
    fill: str,
    outline: str,
    width: int = 3,
    radius: int = 20,
) -> None:
    draw.rounded_rectangle(rect, radius=radius, fill=fill, outline=outline, width=width)


def center(rect: tuple[int, int, int, int]) -> tuple[int, int]:
    return (rect[0] + rect[2]) // 2, (rect[1] + rect[3]) // 2


def point(rect: tuple[int, int, int, int], side: str) -> tuple[int, int]:
    cx, cy = center(rect)
    if side == "left":
        return rect[0], cy
    if side == "right":
        return rect[2], cy
    if side == "top":
        return cx, rect[1]
    if side == "bottom":
        return cx, rect[3]
    return cx, cy


def arrow(
    draw: ImageDraw.ImageDraw,
    start: tuple[int, int],
    end: tuple[int, int],
    color: str = "#64748B",
    width: int = 4,
    label: str | None = None,
) -> None:
    draw.line([start, end], fill=color, width=width)
    angle = atan2(end[1] - start[1], end[0] - start[0])
    length = 18
    spread = pi / 7
    p1 = (end[0] - length * cos(angle - spread), end[1] - length * sin(angle - spread))
    p2 = (end[0] - length * cos(angle + spread), end[1] - length * sin(angle + spread))
    draw.polygon([end, p1, p2], fill=color)
    if label:
        mx = (start[0] + end[0]) // 2
        my = (start[1] + end[1]) // 2
        pad_x = 10
        w = text_width(draw, label, FONT_SMALL)
        draw.rounded_rectangle((mx - w // 2 - pad_x, my - 16, mx + w // 2 + pad_x, my + 18), radius=9, fill="#FFFFFF", outline="#CBD5E1", width=1)
        draw.text((mx - w // 2, my - 12), label, font=FONT_SMALL, fill="#475569")


def path_arrow(
    draw: ImageDraw.ImageDraw,
    points: list[tuple[int, int]],
    color: str = "#64748B",
    width: int = 4,
    label: str | None = None,
) -> None:
    draw.line(points, fill=color, width=width)
    start, end = points[-2], points[-1]
    angle = atan2(end[1] - start[1], end[0] - start[0])
    length = 18
    spread = pi / 7
    p1 = (end[0] - length * cos(angle - spread), end[1] - length * sin(angle - spread))
    p2 = (end[0] - length * cos(angle + spread), end[1] - length * sin(angle + spread))
    draw.polygon([end, p1, p2], fill=color)
    if label:
        mx, my = points[len(points) // 2]
        pad_x = 10
        w = text_width(draw, label, FONT_SMALL)
        draw.rounded_rectangle((mx - w // 2 - pad_x, my - 16, mx + w // 2 + pad_x, my + 18), radius=9, fill="#FFFFFF", outline="#CBD5E1", width=1)
        draw.text((mx - w // 2, my - 12), label, font=FONT_SMALL, fill="#475569")


def draw_wrapped(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int],
    text: str,
    font: ImageFont.ImageFont,
    fill: str,
    max_width: int,
    line_gap: int = 27,
    max_lines: int | None = None,
) -> int:
    x, y = xy
    lines = wrap_text(draw, text, font, max_width)
    if max_lines is not None:
        lines = lines[:max_lines]
    for line in lines:
        draw.text((x, y), line, font=font, fill=fill)
        y += line_gap
    return y


def card(
    draw: ImageDraw.ImageDraw,
    rect: tuple[int, int, int, int],
    title: str,
    body: str,
    palette: dict[str, str],
    tag: str | None = None,
    max_body_lines: int | None = None,
) -> tuple[int, int, int, int]:
    rounded(draw, rect, palette["fill"], palette["outline"], width=3, radius=18)
    draw.rectangle((rect[0], rect[1], rect[0] + 10, rect[3]), fill=palette["accent"])
    x = rect[0] + 25
    y = rect[1] + 18
    if tag:
        draw.rounded_rectangle((x, y, x + 58, y + 34), radius=10, fill=palette["accent"], outline=palette["accent"])
        draw.text((x + 13, y + 5), tag, font=FONT_TAG, fill="#FFFFFF")
        title_x = x + 72
    else:
        title_x = x
    draw.text((title_x, y + 3), title, font=FONT_BOX_TITLE, fill="#111827")
    draw_wrapped(draw, (x, y + 47), body, FONT_BOX, "#374151", rect[2] - rect[0] - 50, max_lines=max_body_lines)
    return rect


def section_label(draw: ImageDraw.ImageDraw, xy: tuple[int, int], label: str, color: str) -> None:
    x, y = xy
    draw.rounded_rectangle((x, y, x + 170, y + 42), radius=14, fill="#FFFFFF", outline=color, width=2)
    draw.text((x + 18, y + 8), label, font=FONT_TAG, fill=color)


PALETTES = {
    "blue": {"fill": "#EFF6FF", "outline": "#60A5FA", "accent": "#2563EB"},
    "purple": {"fill": "#F5F3FF", "outline": "#A78BFA", "accent": "#7C3AED"},
    "green": {"fill": "#ECFDF5", "outline": "#4ADE80", "accent": "#16A34A"},
    "orange": {"fill": "#FFF7ED", "outline": "#FB923C", "accent": "#EA580C"},
    "pink": {"fill": "#FDF2F8", "outline": "#F472B6", "accent": "#DB2777"},
    "slate": {"fill": "#F8FAFC", "outline": "#CBD5E1", "accent": "#64748B"},
}


STAGES = [
    {
        "file": "01_document_parsing.png",
        "title": "01 文档解析流程图",
        "subtitle": "把原始资料转成可追踪、可分块、可入库的标准化文本包。",
        "color": "blue",
        "input": "PDF、Word、TXT、JSON、网页、扫描件、图片/图纸；文件名、来源、版本、上传人、时间。",
        "steps": [
            ("文件登记", "生成 file_id，记录来源、版本、权限与处理状态。"),
            ("格式识别", "识别 PDF/Office/图片/网页，选择解析器与 OCR 策略。"),
            ("版面解析", "提取正文、标题层级、页码、表格、图片和公式位置。"),
            ("OCR 与校正", "扫描件转文字，记录置信度、疑似错字和无法识别区域。"),
            ("结构化输出", "生成 page、block、table、image、raw_text 与 metadata。"),
        ],
        "output": "标准化解析包：pages.json、raw_text.txt、tables.json、images/、parse_log.json。",
        "quality": "页码可追踪；文本覆盖率达标；表格不丢列；低置信 OCR 区域标记待复核。",
        "retention": [
            ("原件快照", "raw_files/ + checksum"),
            ("解析文本", "parsed/pages.json"),
            ("版面对象", "tables.json / images.json"),
            ("质量日志", "parse_log + OCR 置信度"),
        ],
    },
    {
        "file": "02_entity_relation_extraction.png",
        "title": "02 实体关系抽取流程图",
        "subtitle": "从文本片段中抽取设备、部件、故障、原因、措施等实体与关系。",
        "color": "green",
        "input": "解析后的 page/block/chunk 文本，专业 schema，术语词表，故障模式模板。",
        "steps": [
            ("候选实体识别", "抽取设备、部件、参数、故障、现象、原因、措施。"),
            ("关系候选生成", "识别属于、导致、表现为、影响、处理措施等关系。"),
            ("三元组规范化", "统一实体名称、关系类型、单位和同义词。"),
            ("Evidence 绑定", "每个实体/关系绑定原文、页码、chunk_id 与置信度。"),
            ("人工审核", "保留通过、待确认、驳回和修改记录。"),
        ],
        "output": "审核后的 entities.json、relations.json、triples.json、evidence_links.json。",
        "quality": "无证据不入库；实体类型受 schema 约束；关系必须有原文支撑。",
        "retention": [
            ("候选列表", "entity_candidates.json"),
            ("三元组草稿", "triples_draft.json"),
            ("证据链", "evidence_links.json"),
            ("审核记录", "review_decisions.json"),
        ],
    },
    {
        "file": "03_graph_construction.png",
        "title": "03 图谱构建流程图",
        "subtitle": "把审核后的三元组固化成可检索、可计算、可版本管理的知识图谱。",
        "color": "orange",
        "input": "审核后的实体、关系、三元组、证据链接和专业 schema。",
        "steps": [
            ("实体对齐", "合并同义词、别名、大小写和编号差异。"),
            ("关系归一", "清洗关系方向、关系类型和重复边。"),
            ("图结构装配", "生成 node、edge、property、source 字段。"),
            ("属性补全", "补充层级、类别、来源、置信度、时间戳。"),
            ("图谱校验", "检查孤立点、重复边、非法关系和证据缺失。"),
        ],
        "output": "GraphStore、NetworkX/JSON、nodes.csv、edges.csv、graph_stats.json。",
        "quality": "节点边可回溯；图结构符合 schema；每次构建有版本号和差异记录。",
        "retention": [
            ("节点表", "nodes.csv/json"),
            ("边表", "edges.csv/json"),
            ("图统计", "graph_stats.json"),
            ("构建版本", "graph_build_manifest.json"),
        ],
    },
    {
        "file": "04_community_detection.png",
        "title": "04 社区划分流程图",
        "subtitle": "把大图拆成主题更集中的知识社区，为全局检索和摘要提供上下文。",
        "color": "purple",
        "input": "图谱节点、边、权重、关系类型、专业层级和构建版本。",
        "steps": [
            ("图投影与加权", "按关系类型、证据强度、专业层级设置边权。"),
            ("社区算法", "运行 Leiden / Louvain 或规则辅助划分。"),
            ("子图抽取", "为每个 community_id 生成节点、边和核心路径。"),
            ("指标计算", "计算规模、密度、中心节点、模块度和跨社区边。"),
            ("人工调整", "合并过碎社区，拆分过大社区，记录调整依据。"),
        ],
        "output": "community_map.json、community_subgraphs/、community_metrics.csv。",
        "quality": "社区有业务含义；不过度碎片化；核心节点和跨社区关系可解释。",
        "retention": [
            ("社区映射", "community_map.json"),
            ("子图文件", "community_subgraphs/*.json"),
            ("核心节点", "central_nodes.csv"),
            ("划分日志", "community_run_log.json"),
        ],
    },
    {
        "file": "05_topic_summary.png",
        "title": "05 主题摘要流程图",
        "subtitle": "为每个知识社区生成可引用、可复核、可进入 GraphRAG 上下文的主题摘要。",
        "color": "pink",
        "input": "社区子图、核心节点、代表证据片段、专业 schema 和历史摘要。",
        "steps": [
            ("证据筛选", "选取中心节点、关键路径和高质量原文片段。"),
            ("摘要生成", "生成社区主题、关键实体、关系结构和适用场景。"),
            ("标签归类", "标注设备层级、故障模式、参数类别和用途。"),
            ("冲突检查", "识别互相矛盾、证据不足或跨社区重复内容。"),
            ("专家复核", "确认摘要是否专业、完整、可用于问答上下文。"),
        ],
        "output": "community_summary.md/json、topic_tags.json、global_context.json。",
        "quality": "摘要必须带代表证据；不编造不存在的关系；可作为检索上下文复用。",
        "retention": [
            ("摘要草稿", "summary_draft.md"),
            ("代表证据", "summary_evidence.json"),
            ("主题标签", "topic_tags.json"),
            ("复核结论", "summary_review.json"),
        ],
    },
    {
        "file": "06_evidence_retrieval.png",
        "title": "06 证据召回流程图",
        "subtitle": "围绕用户问题召回文本证据、图谱路径和社区摘要，形成回答依据包。",
        "color": "blue",
        "input": "用户问题、历史上下文、向量库、关键词索引、图谱和社区摘要。",
        "steps": [
            ("问题理解", "识别意图、实体、约束、时间范围和答案类型。"),
            ("查询改写", "生成关键词、同义词、实体别名和图谱查询条件。"),
            ("混合召回", "BM25 + Dense + Graph Expansion + Community Summary。"),
            ("重排去重", "按相关性、来源质量、证据新鲜度和覆盖面排序。"),
            ("证据打包", "输出 chunk、页码、原文、相似度、路径和摘要来源。"),
        ],
        "output": "evidence_bundle.json：TopK chunk、图谱路径、社区摘要、引用元数据。",
        "quality": "每条证据可回到原文；召回覆盖问题关键实体；重复片段合并。",
        "retention": [
            ("查询改写", "query_rewrite.json"),
            ("候选证据", "retrieval_candidates.json"),
            ("重排结果", "rerank_scores.csv"),
            ("证据包", "evidence_bundle.json"),
        ],
    },
    {
        "file": "07_reasoning_generation.png",
        "title": "07 推理生成流程图",
        "subtitle": "基于证据包和图谱上下文生成有引用、有边界、有结构的最终回答。",
        "color": "green",
        "input": "用户问题、证据包、图谱路径、社区摘要、回答模板和安全约束。",
        "steps": [
            ("Prompt 装配", "组织问题、证据、引用格式、禁止无证据推断规则。"),
            ("依据链整理", "提炼可展示的推理要点，不保存模型隐藏思考过程。"),
            ("答案生成", "按结论、依据、限制、建议或报告模板输出。"),
            ("引用校验", "检查每个关键结论是否绑定 evidence_id。"),
            ("结构化落盘", "保存最终回答、引用映射、置信度和不确定项。"),
        ],
        "output": "answer.md/json、citation_map.json、uncertainty_flags.json、report.docx。",
        "quality": "无证据则说明不确定；引用不可错位；输出格式稳定可复现。",
        "retention": [
            ("Prompt 快照", "prompt_snapshot.json"),
            ("依据链", "reasoning_basis.json"),
            ("引用映射", "citation_map.json"),
            ("最终输出", "answer.md / report.docx"),
        ],
    },
    {
        "file": "08_result_evaluation.png",
        "title": "08 结果评估流程图",
        "subtitle": "对召回、生成、引用和图谱增强效果做量化评估与失败归因。",
        "color": "slate",
        "input": "问题集、标准答案或人工判断、模型回答、证据包、运行日志。",
        "steps": [
            ("样本归档", "记录问题、答案、证据、模型、参数和运行时间。"),
            ("自动指标", "统计命中率、覆盖率、引用完整性、延迟和成本。"),
            ("质量评审", "LLM Judge + 人工抽检，判断正确性和可用性。"),
            ("失败归因", "定位解析、抽取、图谱、召回、生成或评测问题。"),
            ("改进闭环", "形成 schema、分块、提示词、检索策略的迭代任务。"),
        ],
        "output": "evaluation_report.md/docx、metrics.csv、failure_cases.json、improvement_backlog.md。",
        "quality": "评测样本可复现；失败案例可定位；指标能支撑下一轮优化。",
        "retention": [
            ("评测样本", "eval_cases.jsonl"),
            ("指标表", "metrics.csv"),
            ("失败案例", "failure_cases.json"),
            ("改进清单", "improvement_backlog.md"),
        ],
    },
]


def draw_stage(stage: dict[str, object]) -> Path:
    width, height = 2400, 1500
    image = Image.new("RGB", (width, height), "#F8FAFC")
    draw = ImageDraw.Draw(image)
    palette = PALETTES[str(stage["color"])]

    # Header
    draw.text((80, 48), str(stage["title"]), font=FONT_TITLE, fill="#0F172A")
    draw.text((84, 112), str(stage["subtitle"]), font=FONT_SUBTITLE, fill="#475569")
    draw.rounded_rectangle((80, 165, 2320, 175), radius=5, fill=palette["accent"])

    # Main containers
    rounded(draw, (70, 220, 500, 610), "#FFFFFF", "#CBD5E1", width=2, radius=24)
    rounded(draw, (540, 220, 1860, 760), "#FFFFFF", palette["outline"], width=3, radius=24)
    rounded(draw, (1900, 220, 2330, 610), "#FFFFFF", "#CBD5E1", width=2, radius=24)
    rounded(draw, (70, 820, 2330, 1245), "#FFFFFF", "#CBD5E1", width=2, radius=24)
    rounded(draw, (70, 1285, 2330, 1425), "#FFFFFF", palette["outline"], width=2, radius=24)

    section_label(draw, (105, 250), "输入", palette["accent"])
    draw_wrapped(draw, (115, 315), str(stage["input"]), FONT_BOX, "#334155", 340, line_gap=29)

    section_label(draw, (575, 250), "处理流程", palette["accent"])
    steps = list(stage["steps"])  # type: ignore[arg-type]
    step_rects: list[tuple[int, int, int, int]] = []
    positions = [
        (585, 320, 955, 440),
        (1015, 320, 1385, 440),
        (1445, 320, 1815, 440),
        (800, 560, 1170, 680),
        (1230, 560, 1600, 680),
    ]
    for idx, (title, body) in enumerate(steps):
        rect = positions[idx]
        step_rects.append(card(draw, rect, str(title), str(body), palette, tag=f"S{idx + 1}", max_body_lines=2))

    arrow(draw, point(step_rects[0], "right"), point(step_rects[1], "left"), palette["accent"], 4)
    arrow(draw, point(step_rects[1], "right"), point(step_rects[2], "left"), palette["accent"], 4)
    path_arrow(draw, [point(step_rects[2], "bottom"), (1630, 500), (985, 500), point(step_rects[3], "top")], palette["accent"], 4)
    arrow(draw, point(step_rects[3], "right"), point(step_rects[4], "left"), palette["accent"], 4)

    section_label(draw, (1935, 250), "输出", palette["accent"])
    draw_wrapped(draw, (1945, 315), str(stage["output"]), FONT_BOX, "#334155", 340, line_gap=29)
    draw.text((1945, 500), "质量门禁", font=FONT_BOX_TITLE, fill="#111827")
    draw_wrapped(draw, (1945, 540), str(stage["quality"]), FONT_SMALL, "#475569", 340, line_gap=25)

    arrow(draw, (500, 415), point(step_rects[0], "left"), "#64748B", 4)
    arrow(draw, point(step_rects[-1], "right"), (1900, 415), "#64748B", 4)

    draw.text((110, 855), "中间结果与留存产物", font=FONT_SECTION, fill="#0F172A")
    draw.text((112, 900), "每个环节都要同时保存“可复现输入、处理中间件、质量判断、下一环节输出”。", font=FONT_SUBTITLE, fill="#475569")
    retention = list(stage["retention"])  # type: ignore[arg-type]
    x0s = [110, 660, 1210, 1760]
    for idx, (title, body) in enumerate(retention):
        card(draw, (x0s[idx], 965, x0s[idx] + 460, 1165), str(title), str(body), palette, tag=f"R{idx + 1}", max_body_lines=3)

    draw.text((110, 1316), "留存原则", font=FONT_BOX_TITLE, fill="#111827")
    draw_wrapped(
        draw,
        (255, 1317),
        "文件命名要包含 stage、run_id、版本号和时间；所有结论回到 evidence；失败样本单独归档，进入下一轮改进清单。",
        FONT_BOX,
        "#334155",
        1950,
        line_gap=29,
    )

    out = OUT_DIR / str(stage["file"])
    out.parent.mkdir(parents=True, exist_ok=True)
    image.save(out, quality=95)
    return out


def set_run_font(run, size: int | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color.replace("#", ""))


def add_paragraph(doc: Document, text: str, size: int = 10, bold: bool = False, color: str = "#111827", after: int = 6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def build_docx(image_paths: list[Path]) -> None:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.top_margin = Inches(0.35)
    section.bottom_margin = Inches(0.35)
    section.left_margin = Inches(0.35)
    section.right_margin = Inches(0.35)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("PowerRAG / GraphRAG 关键环节流程图合集")
    set_run_font(run, size=24, bold=True, color="#0F172A")
    add_paragraph(doc, "日期：2026-07-06    用途：整理和留存文档解析、图谱构建、检索生成与评估闭环的中间结果和最终输出。", size=10, color="#475569")
    add_paragraph(doc, "本合集按关键环节拆成 8 张流程图，每张图都包含输入、处理流程、输出、质量门禁和留存产物。可直接作为 WPS 过程结果文档的附图部分。", size=11, color="#111827", after=10)

    for idx, stage in enumerate(STAGES, start=1):
        add_paragraph(doc, f"{idx}. {stage['title']}", size=12, bold=True, color="#0F172A", after=2)
        add_paragraph(doc, str(stage["subtitle"]), size=10, color="#475569", after=6)
    doc.add_page_break()

    for idx, (stage, img_path) in enumerate(zip(STAGES, image_paths), start=1):
        h = doc.add_paragraph()
        h.paragraph_format.space_after = Pt(4)
        run = h.add_run(str(stage["title"]))
        set_run_font(run, size=16, bold=True, color="#0F172A")
        desc = doc.add_paragraph()
        desc.paragraph_format.space_after = Pt(6)
        r = desc.add_run(str(stage["subtitle"]))
        set_run_font(r, size=9, color="#475569")

        pic_p = doc.add_paragraph()
        pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic_run = pic_p.add_run()
        pic_run.add_picture(str(img_path), width=Inches(10.55))
        inline = doc.inline_shapes[-1]._inline
        inline.docPr.set("descr", str(stage["title"]))
        inline.docPr.set("title", str(stage["title"]))

        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(4)
        cr = cap.add_run(f"图 {idx}：{stage['title']}")
        set_run_font(cr, size=9, color="#475569")
        if idx != len(STAGES):
            doc.add_page_break()

    doc.core_properties.title = "PowerRAG / GraphRAG 关键环节流程图合集"
    doc.core_properties.subject = "PowerRAG process flow diagrams"
    doc.core_properties.author = "Codex"
    doc.core_properties.keywords = "PowerRAG, GraphRAG, RAG, flowchart, workflow"
    doc.core_properties.comments = ""
    doc.save(DOCX_OUT)


def build_markdown(image_paths: list[Path]) -> None:
    lines = [
        "# PowerRAG / GraphRAG 关键环节流程图合集",
        "",
        "日期：2026-07-06",
        "",
        "用于系统整理和留存文档解析、实体关系抽取、图谱构建、社区划分、主题摘要、证据召回、推理生成和结果评估等关键环节的中间结果与最终输出。",
        "",
    ]
    for idx, (stage, img_path) in enumerate(zip(STAGES, image_paths), start=1):
        rel = img_path.relative_to(MD_OUT.parent).as_posix()
        lines.extend(
            [
                f"## {idx}. {stage['title']}",
                "",
                str(stage["subtitle"]),
                "",
                f"![{stage['title']}]({rel})",
                "",
            ]
        )
    MD_OUT.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    image_paths = [draw_stage(stage) for stage in STAGES]
    build_docx(image_paths)
    build_markdown(image_paths)
    print(DOCX_OUT)
    print(MD_OUT)
    for path in image_paths:
        print(path)


if __name__ == "__main__":
    main()
