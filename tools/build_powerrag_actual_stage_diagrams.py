from __future__ import annotations

from pathlib import Path
from math import atan2, cos, sin, pi

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "diagrams" / "actual_stage_diagrams"
DOCX_OUT = ROOT / "docs" / "PowerRAG_关键环节流程图合集_实际链路版_2026-07-07.docx"
MD_OUT = ROOT / "docs" / "PowerRAG_关键环节流程图合集_实际链路版_2026-07-07.md"


def load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        Path("C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/simsun.ttc"),
    ]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


F_TITLE = load_font(48, True)
F_SUB = load_font(24)
F_H = load_font(25, True)
F_BOX = load_font(20)
F_SMALL = load_font(17)
F_MONO = load_font(18)


PALETTE = {
    "ink": "#0F172A",
    "muted": "#64748B",
    "line": "#94A3B8",
    "fill": "#FFFFFF",
    "soft": "#F8FAFC",
    "blue": "#2563EB",
    "blue_soft": "#EFF6FF",
    "green": "#16A34A",
    "green_soft": "#ECFDF5",
    "orange": "#EA580C",
    "orange_soft": "#FFF7ED",
    "purple": "#7C3AED",
    "purple_soft": "#F5F3FF",
    "pink": "#DB2777",
    "pink_soft": "#FDF2F8",
    "slate": "#475569",
    "slate_soft": "#F1F5F9",
    "red": "#DC2626",
    "red_soft": "#FEF2F2",
}


def tw(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.ImageFont) -> int:
    if not text:
        return 0
    box = draw.textbbox((0, 0), text, font=font)
    return box[2] - box[0]


def wrap(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.ImageFont, max_width: int) -> list[str]:
    lines: list[str] = []
    current = ""
    break_chars = " ，、；;，。,.：:()/（）-+_"
    for char in text:
        trial = current + char
        if tw(draw, trial, font) <= max_width or not current:
            current = trial
            continue
        idx = max(current.rfind(c) for c in break_chars)
        if idx > 0:
            line = current[: idx + 1].strip()
            rest = current[idx + 1 :].strip()
            if line:
                lines.append(line)
                current = rest + char
                continue
        lines.append(current)
        current = char
    if current:
        lines.append(current)
    return lines


def text_block(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int],
    text: str,
    font: ImageFont.ImageFont = F_BOX,
    fill: str = PALETTE["ink"],
    max_width: int = 300,
    line_gap: int = 28,
    max_lines: int | None = None,
) -> int:
    x, y = xy
    lines = wrap(draw, text, font, max_width)
    if max_lines is not None:
        lines = lines[:max_lines]
    for line in lines:
        draw.text((x, y), line, font=font, fill=fill)
        y += line_gap
    return y


def rounded(draw: ImageDraw.ImageDraw, rect: tuple[int, int, int, int], fill: str, outline: str, width: int = 3, radius: int = 20) -> None:
    draw.rounded_rectangle(rect, radius=radius, fill=fill, outline=outline, width=width)


def box(
    draw: ImageDraw.ImageDraw,
    rect: tuple[int, int, int, int],
    title: str,
    body: str = "",
    fill: str = "#FFFFFF",
    outline: str = "#CBD5E1",
    accent: str | None = None,
    title_font: ImageFont.ImageFont = F_H,
) -> tuple[int, int, int, int]:
    rounded(draw, rect, fill, outline, width=3, radius=18)
    x = rect[0] + 22
    y = rect[1] + 18
    if accent:
        draw.rectangle((rect[0], rect[1], rect[0] + 10, rect[3]), fill=accent)
        x += 8
    draw.text((x, y), title, font=title_font, fill=PALETTE["ink"])
    if body:
        text_block(draw, (x, y + 42), body, F_SMALL, "#334155", rect[2] - rect[0] - 48, line_gap=25)
    return rect


def diamond(draw: ImageDraw.ImageDraw, center: tuple[int, int], w: int, h: int, title: str, body: str, fill: str, outline: str) -> tuple[int, int, int, int]:
    cx, cy = center
    pts = [(cx, cy - h // 2), (cx + w // 2, cy), (cx, cy + h // 2), (cx - w // 2, cy)]
    draw.polygon(pts, fill=fill, outline=outline)
    draw.line([*pts, pts[0]], fill=outline, width=3)
    text_block(draw, (cx - w // 3, cy - 32), title, F_H, PALETTE["ink"], w * 2 // 3, line_gap=28, max_lines=1)
    if body:
        text_block(draw, (cx - w // 3, cy + 4), body, F_SMALL, "#334155", w * 2 // 3, line_gap=23, max_lines=2)
    return (cx - w // 2, cy - h // 2, cx + w // 2, cy + h // 2)


def cylinder(draw: ImageDraw.ImageDraw, rect: tuple[int, int, int, int], title: str, body: str, fill: str, outline: str) -> tuple[int, int, int, int]:
    x1, y1, x2, y2 = rect
    eh = 34
    draw.rectangle((x1, y1 + eh // 2, x2, y2 - eh // 2), fill=fill, outline=outline, width=3)
    draw.ellipse((x1, y1, x2, y1 + eh), fill=fill, outline=outline, width=3)
    draw.ellipse((x1, y2 - eh, x2, y2), fill=fill, outline=outline, width=3)
    draw.text((x1 + 24, y1 + 28), title, font=F_H, fill=PALETTE["ink"])
    text_block(draw, (x1 + 24, y1 + 70), body, F_SMALL, "#334155", x2 - x1 - 48, line_gap=24)
    return rect


def center(rect: tuple[int, int, int, int]) -> tuple[int, int]:
    return ((rect[0] + rect[2]) // 2, (rect[1] + rect[3]) // 2)


def pt(rect: tuple[int, int, int, int], side: str) -> tuple[int, int]:
    cx, cy = center(rect)
    return {
        "left": (rect[0], cy),
        "right": (rect[2], cy),
        "top": (cx, rect[1]),
        "bottom": (cx, rect[3]),
    }.get(side, (cx, cy))


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color: str = "#64748B", width: int = 4, label: str | None = None) -> None:
    draw.line([start, end], fill=color, width=width)
    angle = atan2(end[1] - start[1], end[0] - start[0])
    length = 18
    spread = pi / 7
    p1 = (end[0] - length * cos(angle - spread), end[1] - length * sin(angle - spread))
    p2 = (end[0] - length * cos(angle + spread), end[1] - length * sin(angle + spread))
    draw.polygon([end, p1, p2], fill=color)
    if label:
        mx = (start[0] + end[0]) // 2
        my = (start[1] + end[1]) // 2
        w = tw(draw, label, F_SMALL)
        draw.rounded_rectangle((mx - w // 2 - 10, my - 17, mx + w // 2 + 10, my + 17), radius=9, fill="#FFFFFF", outline="#CBD5E1", width=1)
        draw.text((mx - w // 2, my - 12), label, font=F_SMALL, fill="#475569")


def poly_arrow(draw: ImageDraw.ImageDraw, points: list[tuple[int, int]], color: str = "#64748B", width: int = 4, label: str | None = None) -> None:
    draw.line(points, fill=color, width=width)
    arrow(draw, points[-2], points[-1], color=color, width=width, label=None)
    if label:
        mx, my = points[len(points) // 2]
        w = tw(draw, label, F_SMALL)
        draw.rounded_rectangle((mx - w // 2 - 10, my - 17, mx + w // 2 + 10, my + 17), radius=9, fill="#FFFFFF", outline="#CBD5E1", width=1)
        draw.text((mx - w // 2, my - 12), label, font=F_SMALL, fill="#475569")


def header(draw: ImageDraw.ImageDraw, title: str, subtitle: str, color: str) -> None:
    draw.text((75, 48), title, font=F_TITLE, fill=PALETTE["ink"])
    draw.text((78, 112), subtitle, font=F_SUB, fill=PALETTE["muted"])
    draw.rounded_rectangle((75, 165, 2325, 176), radius=5, fill=color)


def footer(draw: ImageDraw.ImageDraw, text: str) -> None:
    rounded(draw, (75, 1365, 2325, 1445), "#FFFFFF", "#CBD5E1", width=2, radius=18)
    draw.text((105, 1392), "代码依据", font=F_H, fill=PALETTE["ink"])
    text_block(draw, (225, 1395), text, F_SMALL, "#475569", 2000, line_gap=24, max_lines=2)


def new_canvas() -> tuple[Image.Image, ImageDraw.ImageDraw]:
    image = Image.new("RGB", (2400, 1500), PALETTE["soft"])
    return image, ImageDraw.Draw(image)


def save(image: Image.Image, name: str) -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / name
    image.save(path, quality=95)
    return path


def diagram_document_parsing() -> Path:
    image, draw = new_canvas()
    header(draw, "01 文档解析：真实解析路由", "不是固定流程，实际先按文件类型选择 parser_route，再进入分块与质量门禁。", PALETTE["blue"])
    api = box(draw, (90, 245, 390, 355), "/api/process / ingest", "上传文件或目录入库请求", PALETTE["blue_soft"], "#60A5FA", PALETTE["blue"])
    ingest = box(draw, (475, 245, 820, 355), "ingest_source_payloads", "记录 file_summaries，逐文件调用 DocumentIntake", "#FFFFFF", "#CBD5E1")
    decide = diamond(draw, (1030, 305), 290, 170, "classify_document", "按扩展名/内容判断 parser_route", PALETTE["blue_soft"], "#60A5FA")
    arrow(draw, pt(api, "right"), pt(ingest, "left"), PALETTE["blue"])
    arrow(draw, pt(ingest, "right"), pt(decide, "left"), PALETTE["blue"])

    routes = [
        ((1260, 220, 1540, 300), "PDF", "pdf_deepdoc_ready\n必要时标记 OCR"),
        ((1260, 335, 1540, 415), "Office", "office_document\n段落/表格解析"),
        ((1260, 450, 1540, 530), "CSV/TSV", "tabular_document\n表格行转记录"),
        ((1260, 565, 1540, 645), "JSON/JSONL", "structured_json(_lines)"),
        ((1260, 680, 1540, 760), "TXT/MD", "text_document"),
        ((1260, 795, 1540, 875), "unsupported", "直接 failed result"),
    ]
    for rect, title, body in routes:
        r = box(draw, rect, title, body, "#FFFFFF", "#CBD5E1")
        arrow(draw, pt(decide, "right"), pt(r, "left"), "#64748B")

    records = box(draw, (1650, 390, 1965, 505), "SourceRecord", "原文、页码、label、metadata", PALETTE["blue_soft"], "#60A5FA", PALETTE["blue"])
    chunks = box(draw, (1650, 620, 1965, 750), "chunk_records", "Title 触发新 chunk；split_text_with_overlap；稳定 chunk_id", PALETTE["blue_soft"], "#60A5FA", PALETTE["blue"])
    quality = box(draw, (2050, 500, 2300, 635), "quality_report", "chunk_count、缺失 metadata、门禁状态", "#FFFFFF", "#CBD5E1")
    for rect, _, _ in routes[:5]:
        poly_arrow(draw, [pt(rect, "right"), (1600, center(rect)[1]), (1600, 448), pt(records, "left")], "#64748B")
    arrow(draw, pt(records, "bottom"), pt(chunks, "top"), PALETTE["blue"], label="parsed")
    arrow(draw, pt(chunks, "right"), pt(quality, "left"), PALETTE["blue"])
    upsert = cylinder(draw, (2050, 745, 2300, 905), "Chroma upsert", "ids / documents / metadatas\n写入 collection", PALETTE["blue_soft"], "#60A5FA")
    arrow(draw, pt(chunks, "right"), pt(upsert, "left"), PALETTE["blue"])
    manifest = box(draw, (145, 760, 450, 900), "实际留存", "upload manifest\nparse_file_summary\nchunk_preview\noperation log", "#FFFFFF", "#CBD5E1")
    poly_arrow(draw, [pt(ingest, "bottom"), (650, 705), pt(manifest, "right")], "#64748B")
    footer(draw, "data_pipeline/document_intake.py: classify_document/run_document_intake；pipeline.py: ingest_source_payloads；chunking.py: chunk_records；api.py: /api/process, /api/ingest")
    return save(image, "01_document_parsing_actual.png")


def diagram_entity_extraction() -> Path:
    image, draw = new_canvas()
    header(draw, "02 实体关系抽取：LLM 抽取 + Evidence 校验链", "实际实现围绕 chunk、schema、LLM JSON、normalize 和 validate，不是普通 NLP 流程图。", PALETTE["green"])
    chunks = cylinder(draw, (90, 300, 360, 455), "chunks 输入", "read_chunks(input_path)\nchunk_id/source/page/text", PALETTE["green_soft"], "#4ADE80")
    schema = box(draw, (90, 585, 360, 710), "schema", "relation_types\n专业关系白名单", "#FFFFFF", "#CBD5E1")
    prompt = box(draw, (510, 320, 840, 470), "_build_prompt", "要求输出 JSON：triples[{subject, relation, object, evidence, source, page, valid_time}]", PALETTE["green_soft"], "#4ADE80", PALETTE["green"])
    llm = box(draw, (970, 320, 1250, 470), "_call_llm", "OpenAICompatibleClient 或外部 LLM client", "#FFFFFF", "#CBD5E1")
    raw = box(draw, (1380, 320, 1660, 470), "raw_response", "模型返回 JSON 文本", "#FFFFFF", "#CBD5E1")
    parse = box(draw, (1790, 300, 2120, 455), "_parse_triples", "必须包含 triples list；否则进入 error", PALETTE["green_soft"], "#4ADE80", PALETTE["green"])
    normalize = box(draw, (510, 760, 840, 910), "_normalize_triple", "补 source/page/chunk_id；字段转文本；保留 evidence", "#FFFFFF", "#CBD5E1")
    validate = diamond(draw, (1110, 835), 300, 180, "_validate_triple", "必填字段 + relation_types", PALETTE["green_soft"], "#4ADE80")
    ok = cylinder(draw, (1430, 735, 1725, 900), "triples 输出", "subject/relation/object\nevidence/source/page", PALETTE["green_soft"], "#4ADE80")
    err = box(draw, (1430, 955, 1725, 1075), "errors", "无 evidence、关系非法、JSON 解析失败", PALETTE["red_soft"], "#FCA5A5", PALETTE["red"])
    review = box(draw, (1890, 790, 2220, 940), "人工复核入口", "manual_review.csv\n通过/修改/驳回", "#FFFFFF", "#CBD5E1")
    arrow(draw, pt(chunks, "right"), pt(prompt, "left"), PALETTE["green"])
    arrow(draw, pt(schema, "right"), (510, 395), "#64748B")
    arrow(draw, pt(prompt, "right"), pt(llm, "left"), PALETTE["green"])
    arrow(draw, pt(llm, "right"), pt(raw, "left"), PALETTE["green"])
    arrow(draw, pt(raw, "right"), pt(parse, "left"), PALETTE["green"])
    poly_arrow(draw, [pt(parse, "bottom"), (1955, 675), (675, 675), pt(normalize, "top")], PALETTE["green"])
    arrow(draw, pt(normalize, "right"), pt(validate, "left"), PALETTE["green"])
    arrow(draw, pt(validate, "right"), pt(ok, "left"), PALETTE["green"], label="valid")
    poly_arrow(draw, [pt(validate, "bottom"), (1110, 1030), pt(err, "left")], PALETTE["red"], label="invalid")
    arrow(draw, pt(ok, "right"), pt(review, "left"), "#64748B")
    footer(draw, "kg_pipeline/llm_extraction/pipeline.py: read_chunks, run_extraction, _build_prompt, _parse_triples, _normalize_triple, _validate_triple；kg_pipeline/poc/manual_review.csv")
    return save(image, "02_entity_relation_extraction_actual.png")


def diagram_graph_construction() -> Path:
    image, draw = new_canvas()
    header(draw, "03 图谱构建：GraphStore 实际数据结构", "实际不是画概念知识图谱，而是把 triples/links 规范化后写进 SQLite GraphStore。", PALETTE["orange"])
    input_box = box(draw, (90, 265, 410, 430), "输入形态", "triples.json 或 graph_links\nnormalize_kg_payload", PALETTE["orange_soft"], "#FDBA74", PALETTE["orange"])
    edge = box(draw, (535, 265, 875, 430), "GraphEdgeRecord", "subject / predicate / object_name\nsource/page/evidence/confidence", "#FFFFFF", "#CBD5E1")
    store = cylinder(draw, (1035, 245, 1360, 450), "GraphStore(db_path)", "SQLite 初始化与写入\ninitialize(reset)", PALETTE["orange_soft"], "#FDBA74")
    arrow(draw, pt(input_box, "right"), pt(edge, "left"), PALETTE["orange"])
    arrow(draw, pt(edge, "right"), pt(store, "left"), PALETTE["orange"])

    tables = [
        ((250, 640, 620, 810), "nodes", "id, name, entity_type\n唯一实体节点"),
        ((760, 640, 1130, 810), "edges", "subject_id, predicate,\nobject_id, source, page,\nconfidence, metadata_json"),
        ((1270, 640, 1640, 810), "evidence", "edge_id, triple_id,\ntext, source, page"),
        ((520, 950, 890, 1120), "communities", "community_id, level,\nnode_id"),
        ((1030, 950, 1400, 1120), "community_summaries", "community_id, title,\nsummary, metadata_json"),
    ]
    for rect, title, body in tables:
        cylinder(draw, rect, title, body, "#FFFFFF", "#CBD5E1")
    for rect, _, _ in tables[:3]:
        arrow(draw, pt(store, "bottom"), pt(rect, "top"), "#64748B")
    arrow(draw, pt(tables[0][0], "right"), pt(tables[1][0], "left"), PALETTE["orange"], label="subject/object")
    arrow(draw, pt(tables[1][0], "right"), pt(tables[2][0], "left"), PALETTE["orange"], label="edge_id")
    arrow(draw, pt(tables[3][0], "right"), pt(tables[4][0], "left"), PALETTE["orange"], label="summary")

    export = box(draw, (1760, 315, 2210, 520), "出口", "summary()\nexport_graph()\nwrite_neo4j_cypher()\n/api/graphrag/export", PALETTE["orange_soft"], "#FDBA74", PALETTE["orange"])
    arrow(draw, pt(store, "right"), pt(export, "left"), PALETTE["orange"])
    routes = box(draw, (1680, 810, 2260, 1035), "API 侧真实入口", "routes_graphrag.py\n/import：浏览器图谱快照导入\n/reset：删除 SQLite GraphStore\n/stats：节点/边/evidence 统计", "#FFFFFF", "#CBD5E1")
    arrow(draw, pt(routes, "left"), pt(store, "right"), "#64748B")
    footer(draw, "storage_layer/graph_store.py: GraphStore.initialize/import_edges/export_graph/store_communities/store_community_summaries；routes_graphrag.py: /import, /reset, /stats, /export")
    return save(image, "03_graph_construction_actual.png")


def diagram_community_detection() -> Path:
    image, draw = new_canvas()
    header(draw, "04 社区划分：Leiden 聚类与 fallback", "实际从 GraphStore 边表构图，优先跑 Leiden；缺依赖或导入场景会走连通分量 fallback。", PALETTE["purple"])
    graphstore = cylinder(draw, (95, 330, 395, 520), "GraphStore", "get_all_edges_as_tuples()\nnodes + edges", PALETTE["purple_soft"], "#A78BFA")
    nx = box(draw, (555, 300, 875, 440), "build_networkx_graph", "把 SQLite 边表变为 NetworkX Graph", "#FFFFFF", "#CBD5E1")
    alg = diamond(draw, (1125, 370), 340, 180, "依赖可用？", "networkx + igraph + leidenalg", PALETTE["purple_soft"], "#A78BFA")
    leiden = box(draw, (1390, 255, 1760, 405), "run_leiden_detection", "resolution=0.7\nCommunityAssignment(C0/C1/...)", PALETTE["purple_soft"], "#A78BFA", PALETTE["purple"])
    fallback = box(draw, (1390, 510, 1760, 660), "fallback components", "导入图谱快照时按连通分量生成 assignments", "#FFFFFF", "#CBD5E1")
    store = cylinder(draw, (1910, 360, 2220, 545), "store_communities", "communities 表\ncommunity_id, level, node_id", PALETTE["purple_soft"], "#A78BFA")
    arrow(draw, pt(graphstore, "right"), pt(nx, "left"), PALETTE["purple"])
    arrow(draw, pt(nx, "right"), pt(alg, "left"), PALETTE["purple"])
    arrow(draw, pt(alg, "right"), pt(leiden, "left"), PALETTE["purple"], label="yes")
    poly_arrow(draw, [pt(alg, "bottom"), (1125, 590), pt(fallback, "left")], "#64748B", label="fallback")
    arrow(draw, pt(leiden, "right"), pt(store, "left"), PALETTE["purple"])
    arrow(draw, pt(fallback, "right"), pt(store, "left"), "#64748B")

    # Cluster illustration
    rounded(draw, (180, 790, 2200, 1245), "#FFFFFF", "#CBD5E1", width=2, radius=22)
    draw.text((220, 825), "划分结果不是流程步数，而是图结构上的社区归属", font=F_H, fill=PALETTE["ink"])
    clusters = [
        ((520, 1010), "#C4B5FD", "C0"),
        ((1030, 1010), "#A7F3D0", "C1"),
        ((1540, 1010), "#FDBA74", "C2"),
    ]
    for (cx, cy), color, label in clusters:
        draw.ellipse((cx - 150, cy - 105, cx + 150, cy + 105), fill=color, outline="#475569", width=3)
        draw.text((cx - 22, cy - 16), label, font=F_H, fill=PALETTE["ink"])
        for dx, dy in [(-70, -35), (35, -50), (70, 38), (-20, 50)]:
            draw.ellipse((cx + dx - 18, cy + dy - 18, cx + dx + 18, cy + dy + 18), fill="#FFFFFF", outline="#475569", width=2)
    arrow(draw, (670, 1010), (880, 1010), "#94A3B8")
    arrow(draw, (1180, 1010), (1390, 1010), "#94A3B8")
    footer(draw, "kg_pipeline/community_detection.py: build_networkx_graph, run_leiden_detection, run_hierarchical_detection；routes_graphrag.py: _fallback_component_assignments, /community/detect")
    return save(image, "04_community_detection_actual.png")


def diagram_topic_summary() -> Path:
    image, draw = new_canvas()
    header(draw, "05 主题摘要：按社区 Map，再绑定句子级证据", "实际不是单个摘要框，而是遍历 communities，对每个社区取实体/边、调用 LLM、再做 sentence evidence。", PALETTE["pink"])
    start = cylinder(draw, (95, 300, 380, 480), "GraphStore", "get_communities(level)\ncommunity_id/member_count", PALETTE["pink_soft"], "#F472B6")
    filter_box = diamond(draw, (620, 390), 310, 170, "size >= min?", "默认 min_community_size=2", PALETTE["pink_soft"], "#F472B6")
    entities = box(draw, (845, 250, 1165, 385), "get_community_entities", "社区成员实体列表", "#FFFFFF", "#CBD5E1")
    edges = box(draw, (845, 465, 1165, 600), "get_community_edges", "社区内部边 + evidence", "#FFFFFF", "#CBD5E1")
    prompt = box(draw, (1310, 335, 1635, 500), "build_community_prompt", "实体、关系、数量进入 prompt\n要求 JSON title/summary", PALETTE["pink_soft"], "#F472B6", PALETTE["pink"])
    llm = box(draw, (1765, 335, 2050, 500), "LLM generate", "解析 JSON；失败时 raw text fallback", "#FFFFFF", "#CBD5E1")
    evidence = box(draw, (1290, 760, 1645, 930), "build_summary_sentence_evidence", "把摘要句子匹配到 community edges；无匹配则回退到社区边集", PALETTE["pink_soft"], "#F472B6", PALETTE["pink"])
    store = cylinder(draw, (1780, 735, 2135, 955), "store_community_summaries", "community_summaries 表\nmetadata: sentence_evidence", PALETTE["pink_soft"], "#F472B6")
    arrow(draw, pt(start, "right"), pt(filter_box, "left"), PALETTE["pink"])
    arrow(draw, pt(filter_box, "right"), pt(entities, "left"), PALETTE["pink"], label="yes")
    arrow(draw, pt(filter_box, "right"), pt(edges, "left"), PALETTE["pink"])
    arrow(draw, pt(entities, "right"), pt(prompt, "left"), PALETTE["pink"])
    arrow(draw, pt(edges, "right"), pt(prompt, "left"), PALETTE["pink"])
    arrow(draw, pt(prompt, "right"), pt(llm, "left"), PALETTE["pink"])
    poly_arrow(draw, [pt(llm, "bottom"), (1900, 650), pt(evidence, "right")], PALETTE["pink"])
    arrow(draw, pt(evidence, "right"), pt(store, "left"), PALETTE["pink"])
    skip = box(draw, (510, 680, 790, 805), "skip list", "小社区跳过，写入 errors/日志", "#FFFFFF", "#CBD5E1")
    poly_arrow(draw, [pt(filter_box, "bottom"), (620, 610), pt(skip, "top")], "#64748B", label="no")
    rounded(draw, (145, 1030, 2245, 1245), "#FFFFFF", "#CBD5E1", width=2, radius=22)
    draw.text((185, 1068), "输出不是单一摘要", font=F_H, fill=PALETTE["ink"])
    text_block(draw, (185, 1115), "CommunitySummaryResult 同时包含 summaries、errors 和 stored_count；每条 summary 包含 community_id、title、summary、entity_count、edge_count、metadata。", F_BOX, "#334155", 1950, line_gap=30)
    footer(draw, "kg_pipeline/community_summary.py: summarize_communities, build_community_prompt, build_summary_sentence_evidence；GraphStore.get_community_entities/get_community_edges/store_community_summaries")
    return save(image, "05_topic_summary_actual.png")


def diagram_evidence_retrieval() -> Path:
    image, draw = new_canvas()
    header(draw, "06 证据召回：多路召回 + RRF + 邻近扩展 + 重排", "实际 query_collection 不是单一路径，会结合策略、改写、过滤、向量、关键词、图谱和 reranker。", PALETTE["blue"])
    query = box(draw, (90, 270, 360, 405), "query_collection", "query_text, top_k,\nfilters, policy_settings", PALETTE["blue_soft"], "#60A5FA", PALETTE["blue"])
    rewrite = box(draw, (500, 215, 820, 350), "TemplateQueryRewriter", "生成 rewritten_queries；保留 original", "#FFFFFF", "#CBD5E1")
    filters = box(draw, (500, 430, 820, 565), "_auto_filters_from_query", "年份/来源等 metadata filter", "#FFFFFF", "#CBD5E1")
    arrow(draw, pt(query, "right"), pt(rewrite, "left"), PALETTE["blue"])
    arrow(draw, pt(query, "right"), pt(filters, "left"), "#64748B")
    vector = box(draw, (1000, 190, 1320, 325), "Chroma 向量召回", "_ChromaCollectionRetriever\ncollection.query()", PALETTE["blue_soft"], "#60A5FA", PALETTE["blue"])
    keyword = box(draw, (1000, 380, 1320, 515), "关键词召回", "KeywordRetriever\nmetadata-aware sparse", PALETTE["orange_soft"], "#FDBA74", PALETTE["orange"])
    graph = box(draw, (1000, 555, 1320, 750), "图谱召回", "SQLiteGraphRetriever\nentity match\nPPR/BFS\ncommunity summary", PALETTE["purple_soft"], "#A78BFA", PALETTE["purple"])
    arrow(draw, pt(rewrite, "right"), pt(vector, "left"), PALETTE["blue"])
    arrow(draw, pt(rewrite, "right"), pt(keyword, "left"), PALETTE["orange"])
    arrow(draw, pt(filters, "right"), pt(vector, "left"), "#64748B")
    arrow(draw, pt(filters, "right"), pt(keyword, "left"), "#64748B")
    arrow(draw, pt(filters, "right"), pt(graph, "left"), "#64748B")
    fuse = diamond(draw, (1590, 445), 340, 210, "HybridRetriever", "RRF 融合 + component weights", PALETTE["blue_soft"], "#60A5FA")
    for r, c in [(vector, PALETTE["blue"]), (keyword, PALETTE["orange"]), (graph, PALETTE["purple"])]:
        arrow(draw, pt(r, "right"), pt(fuse, "left"), c)
    neighbor = box(draw, (1805, 225, 2190, 360), "neighbor/backfill", "source-local 邻近 chunk\n补足同源上下文", "#FFFFFF", "#CBD5E1")
    rerank = box(draw, (1805, 500, 2190, 635), "reranker", "cross-encoder 可选；失败写 diagnostics", "#FFFFFF", "#CBD5E1")
    result = cylinder(draw, (1805, 765, 2190, 960), "evidence bundle", "API result：text, source,\npage, chunk_id, score,\ndiagnostics", PALETTE["blue_soft"], "#60A5FA")
    arrow(draw, pt(fuse, "right"), pt(neighbor, "left"), PALETTE["blue"])
    arrow(draw, pt(neighbor, "bottom"), pt(rerank, "top"), PALETTE["blue"])
    arrow(draw, pt(rerank, "bottom"), pt(result, "top"), PALETTE["blue"])
    rounded(draw, (160, 1025, 1540, 1235), "#FFFFFF", "#CBD5E1", width=2, radius=22)
    draw.text((200, 1060), "实际诊断字段", font=F_H, fill=PALETTE["ink"])
    text_block(draw, (200, 1110), "retriever weights、rewritten_queries、reranker_name/error、no_answer_reason、auto_filters、result scores 会进入 diagnostics，方便评估和调参。", F_BOX, "#334155", 1250, line_gap=30)
    footer(draw, "pipeline.py: query_collection, _ChromaCollectionRetriever, _auto_filters_from_query, _TemplateQueryRewriter；retrieval_engine/hybrid.py: HybridRetriever RRF/reranker；retrieval_engine/graph.py: SQLiteGraphRetriever")
    return save(image, "06_evidence_retrieval_actual.png")


def diagram_reasoning_generation() -> Path:
    image, draw = new_canvas()
    header(draw, "07 推理生成：GraphRAG QA 编排状态机", "实际由 GraphRagQAOrchestrator 根据路由选择 text、graph、global context，再做证据约束生成。", PALETTE["green"])
    q = box(draw, (90, 305, 360, 430), "question", "用户问题 + top_k + context_only", PALETTE["green_soft"], "#4ADE80", PALETTE["green"])
    route = diamond(draw, (580, 365), 300, 180, "route_query", "VECTOR_ONLY / LOCAL / GLOBAL", PALETTE["green_soft"], "#4ADE80")
    text = box(draw, (840, 190, 1165, 320), "text_retriever", "Chroma/Hybrid 返回 text evidence [T]", PALETTE["blue_soft"], "#60A5FA", PALETTE["blue"])
    graph = box(draw, (840, 390, 1165, 520), "graph_retriever", "entity-local triples / graph evidence [G]", PALETTE["purple_soft"], "#A78BFA", PALETTE["purple"])
    global_ctx = box(draw, (840, 590, 1165, 735), "global_searcher", "community summaries\nsource_evidence", PALETTE["pink_soft"], "#F472B6", PALETTE["pink"])
    context = box(draw, (1330, 345, 1665, 520), "build_context", "Text retrieval evidence\nGraph retrieval evidence\nGlobal community context", "#FFFFFF", "#CBD5E1")
    prompt = box(draw, (1830, 275, 2200, 430), "build_default_prompt", "Use only retrieved evidence；不足则说明；引用 [T]/[G]", PALETTE["green_soft"], "#4ADE80", PALETTE["green"])
    llm = box(draw, (1830, 555, 2200, 690), "LLM generate", "answer markdown", "#FFFFFF", "#CBD5E1")
    guard = diamond(draw, (1510, 875), 340, 190, "hallucination_guard", "发现无证据 claim？", PALETTE["green_soft"], "#4ADE80")
    final = cylinder(draw, (1830, 815, 2200, 1025), "GraphRagQAResult", "answer\nmode\ncitations/evidence\ncontext", PALETTE["green_soft"], "#4ADE80")
    arrow(draw, pt(q, "right"), pt(route, "left"), PALETTE["green"])
    arrow(draw, pt(route, "right"), pt(text, "left"), PALETTE["blue"])
    arrow(draw, pt(route, "right"), pt(graph, "left"), PALETTE["purple"])
    arrow(draw, pt(route, "right"), pt(global_ctx, "left"), PALETTE["pink"])
    for r in [text, graph, global_ctx]:
        arrow(draw, pt(r, "right"), pt(context, "left"), "#64748B")
    arrow(draw, pt(context, "right"), pt(prompt, "left"), PALETTE["green"])
    arrow(draw, pt(prompt, "bottom"), pt(llm, "top"), PALETTE["green"])
    arrow(draw, pt(llm, "bottom"), pt(guard, "right"), PALETTE["green"])
    arrow(draw, pt(guard, "right"), pt(final, "left"), PALETTE["green"], label="pass")
    poly_arrow(draw, [pt(guard, "top"), (1510, 760), (1990, 760), pt(prompt, "bottom")], PALETTE["red"], label="correction prompt")
    rounded(draw, (145, 1020, 1130, 1235), "#FFFFFF", "#CBD5E1", width=2, radius=22)
    draw.text((185, 1055), "实际关键约束", font=F_H, fill=PALETTE["ink"])
    text_block(draw, (185, 1105), "prompt 明确要求只使用检索证据；citations 来自 normalized text evidence 和 graph evidence；global context 会拆成 community source evidence。", F_BOX, "#334155", 890, line_gap=30)
    footer(draw, "rag_orchestrator/graphrag_qa.py: GraphRagQAOrchestrator.answer, build_context, build_default_prompt, _normalize_text_evidence, _normalize_graph_evidence, hallucination_guard")
    return save(image, "07_reasoning_generation_actual.png")


def diagram_result_evaluation() -> Path:
    image, draw = new_canvas()
    header(draw, "08 结果评估：指标门禁与策略回写闭环", "实际评估不是最后打分，而是用 threshold 触发 failure_cases 和 retrieval policy 建议。", PALETTE["slate"])
    cases = cylinder(draw, (95, 280, 395, 450), "RAGEvaluationCase", "question\nexpected_evidence_keywords\nreference_answer", PALETTE["slate_soft"], "#CBD5E1")
    harness = box(draw, (560, 285, 870, 445), "RAGEvaluationHarness.run", "逐 case 调 query_fn 或 retriever", "#FFFFFF", "#CBD5E1")
    normalize = box(draw, (1020, 285, 1325, 445), "_normalize_rag_output", "hits/contexts/documents/evidence 统一成评估输入", "#FFFFFF", "#CBD5E1")
    metrics = box(draw, (1480, 215, 1855, 515), "evaluate_records / metrics", "retrieval recall\ncitation missing\nhallucination risk\nfaithfulness / relevancy\ncontext recall / completeness", PALETTE["slate_soft"], "#CBD5E1", PALETTE["slate"])
    gate = diamond(draw, (2045, 365), 320, 190, "Thresholds.check", "达到门禁？", "#FFFFFF", "#CBD5E1")
    arrow(draw, pt(cases, "right"), pt(harness, "left"), PALETTE["slate"])
    arrow(draw, pt(harness, "right"), pt(normalize, "left"), PALETTE["slate"])
    arrow(draw, pt(normalize, "right"), pt(metrics, "left"), PALETTE["slate"])
    arrow(draw, pt(metrics, "right"), pt(gate, "left"), PALETTE["slate"])
    report = cylinder(draw, (1415, 780, 1745, 990), "RAGEvaluationReport", "to_markdown()\nmetrics + gate_status\nfailure_cases", PALETTE["slate_soft"], "#CBD5E1")
    policy = box(draw, (1855, 780, 2225, 990), "默认策略建议", "query_rewrite\nreranker\nno-answer gating\ngraph_retriever 条件启用", "#FFFFFF", "#CBD5E1")
    failure = box(draw, (925, 780, 1285, 990), "_select_failure_cases", "把低召回、缺引用、幻觉风险样本沉淀", PALETTE["red_soft"], "#FCA5A5", PALETTE["red"])
    arrow(draw, pt(gate, "bottom"), pt(report, "right"), PALETTE["slate"])
    arrow(draw, pt(report, "right"), pt(policy, "left"), PALETTE["slate"])
    arrow(draw, pt(report, "left"), pt(failure, "right"), PALETTE["red"])
    poly_arrow(draw, [pt(policy, "top"), (2040, 660), (710, 660), pt(harness, "bottom")], PALETTE["orange"], label="下一轮检索策略")
    rounded(draw, (135, 1065, 800, 1240), "#FFFFFF", "#CBD5E1", width=2, radius=22)
    draw.text((175, 1100), "输出文件", font=F_H, fill=PALETTE["ink"])
    text_block(draw, (175, 1150), "report.md/json、metrics、failure_cases、policy 建议；用于回写检索、分块、引用门禁和 GraphRAG 策略。", F_BOX, "#334155", 580, line_gap=30)
    footer(draw, "evaluation/harness.py: RAGEvaluationHarness, EvaluationThresholds, RAGEvaluationReport, _select_failure_cases, _build_retrieval_default_policy；evaluation/metrics.py: evaluate_single")
    return save(image, "08_result_evaluation_actual.png")


DIAGRAMS = [
    ("文档解析", diagram_document_parsing),
    ("实体关系抽取", diagram_entity_extraction),
    ("图谱构建", diagram_graph_construction),
    ("社区划分", diagram_community_detection),
    ("主题摘要", diagram_topic_summary),
    ("证据召回", diagram_evidence_retrieval),
    ("推理生成", diagram_reasoning_generation),
    ("结果评估", diagram_result_evaluation),
]


def set_run_font(run, size: int | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color.replace("#", ""))


def para(doc: Document, text: str, *, size: int = 10, bold: bool = False, color: str = "#111827", after: int = 6) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color)


def build_docx(paths: list[Path]) -> None:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.top_margin = Inches(0.35)
    section.bottom_margin = Inches(0.35)
    section.left_margin = Inches(0.35)
    section.right_margin = Inches(0.35)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    r = title.add_run("PowerRAG / GraphRAG 关键环节流程图合集：实际链路版")
    set_run_font(r, size=23, bold=True, color="#0F172A")
    para(doc, "日期：2026-07-07    说明：本版按代码中的真实模块、数据结构和接口关系绘制，不使用统一模板画法。", size=10, color="#475569")
    para(doc, "图型分别采用：解析路由图、LLM 抽取校验链、GraphStore 数据结构图、Leiden 社区聚类图、社区摘要 map 图、多路召回漏斗、GraphRAG QA 状态机、评估门禁闭环。", size=10, color="#111827", after=10)
    for i, (name, _) in enumerate(DIAGRAMS, 1):
        para(doc, f"{i}. {name}", size=12, bold=True, color="#0F172A", after=2)
    doc.add_page_break()
    for i, ((name, _), path) in enumerate(zip(DIAGRAMS, paths), 1):
        para(doc, f"{i}. {name}", size=16, bold=True, color="#0F172A", after=4)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(path), width=Inches(10.55))
        inline = doc.inline_shapes[-1]._inline
        inline.docPr.set("title", name)
        inline.docPr.set("descr", f"{name} actual implementation diagram")
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(4)
        cr = cap.add_run(f"图 {i}：{name}（实际链路版）")
        set_run_font(cr, size=9, color="#475569")
        if i != len(paths):
            doc.add_page_break()
    doc.core_properties.title = "PowerRAG / GraphRAG 关键环节流程图合集：实际链路版"
    doc.core_properties.subject = "Actual implementation diagrams for PowerRAG"
    doc.core_properties.author = "Codex"
    doc.core_properties.keywords = "PowerRAG, GraphRAG, actual workflow, diagram"
    doc.core_properties.comments = ""
    doc.save(DOCX_OUT)


def build_markdown(paths: list[Path]) -> None:
    lines = [
        "# PowerRAG / GraphRAG 关键环节流程图合集：实际链路版",
        "",
        "日期：2026-07-07",
        "",
        "本版按代码中的真实模块、数据结构和接口关系绘制，不使用统一模板画法。",
        "",
    ]
    for i, ((name, _), path) in enumerate(zip(DIAGRAMS, paths), 1):
        rel = path.relative_to(MD_OUT.parent).as_posix()
        lines.extend([f"## {i}. {name}", "", f"![{name}]({rel})", ""])
    MD_OUT.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    paths = [builder() for _, builder in DIAGRAMS]
    build_docx(paths)
    build_markdown(paths)
    print(DOCX_OUT)
    print(MD_OUT)
    for path in paths:
        print(path)


if __name__ == "__main__":
    main()
