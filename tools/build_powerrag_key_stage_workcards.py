from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "diagrams" / "key_stage_workcards"
DOCX_OUT = ROOT / "docs" / "PowerRAG_关键环节流程图合集_工作卡版_2026-07-07.docx"
MD_OUT = ROOT / "docs" / "PowerRAG_关键环节流程图合集_工作卡版_2026-07-07.md"


def load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        Path("C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/simsun.ttc"),
    ]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


FONT_TITLE = load_font(50, True)
FONT_SUBTITLE = load_font(25)
FONT_H2 = load_font(27, True)
FONT_CARD_TITLE = load_font(24, True)
FONT_BODY = load_font(20)
FONT_SMALL = load_font(18)
FONT_CHIP = load_font(20, True)


def text_width(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.ImageFont) -> int:
    if not text:
        return 0
    box = draw.textbbox((0, 0), text, font=font)
    return box[2] - box[0]


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.ImageFont, max_width: int) -> list[str]:
    lines: list[str] = []
    current = ""
    break_chars = " ，、；;，。,.：:()/（）-+"
    for char in text:
        trial = current + char
        if text_width(draw, trial, font) <= max_width or not current:
            current = trial
        else:
            break_at = max(current.rfind(c) for c in break_chars)
            if break_at > 0:
                line = current[: break_at + 1].strip()
                rest = current[break_at + 1 :].strip()
                if line:
                    lines.append(line)
                    current = rest + char
                    continue
            lines.append(current)
            current = char
    if current:
        lines.append(current)
    return lines


def rounded(draw: ImageDraw.ImageDraw, rect: tuple[int, int, int, int], fill: str, outline: str, width: int = 3, radius: int = 22) -> None:
    draw.rounded_rectangle(rect, radius=radius, fill=fill, outline=outline, width=width)


def draw_wrapped(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int],
    text: str,
    font: ImageFont.ImageFont,
    fill: str,
    max_width: int,
    line_gap: int = 29,
    max_lines: int | None = None,
) -> int:
    x, y = xy
    lines = wrap_text(draw, text, font, max_width)
    if max_lines is not None:
        lines = lines[:max_lines]
    for line in lines:
        draw.text((x, y), line, font=font, fill=fill)
        y += line_gap
    return y


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color: str, width: int = 4) -> None:
    x1, y1 = start
    x2, y2 = end
    draw.line((x1, y1, x2, y2), fill=color, width=width)
    if abs(x2 - x1) >= abs(y2 - y1):
        sign = 1 if x2 >= x1 else -1
        pts = [(x2, y2), (x2 - sign * 20, y2 - 10), (x2 - sign * 20, y2 + 10)]
    else:
        sign = 1 if y2 >= y1 else -1
        pts = [(x2, y2), (x2 - 10, y2 - sign * 20), (x2 + 10, y2 - sign * 20)]
    draw.polygon(pts, fill=color)


def panel(draw: ImageDraw.ImageDraw, rect: tuple[int, int, int, int], title: str, palette: dict[str, str]) -> None:
    rounded(draw, rect, "#FFFFFF", palette["outline"], width=2, radius=24)
    draw.rectangle((rect[0], rect[1], rect[0] + 10, rect[3]), fill=palette["accent"])
    draw.text((rect[0] + 30, rect[1] + 24), title, font=FONT_H2, fill="#0F172A")


def bullet_list(draw: ImageDraw.ImageDraw, items: list[str], x: int, y: int, max_width: int, max_lines_each: int = 2) -> int:
    for item in items:
        draw.ellipse((x, y + 9, x + 10, y + 19), fill="#64748B")
        y = draw_wrapped(draw, (x + 24, y), item, FONT_BODY, "#334155", max_width - 24, line_gap=28, max_lines=max_lines_each)
        y += 10
    return y


def chip(draw: ImageDraw.ImageDraw, rect: tuple[int, int, int, int], text: str, palette: dict[str, str]) -> None:
    rounded(draw, rect, palette["fill"], palette["outline"], width=3, radius=18)
    draw.text((rect[0] + 24, rect[1] + 18), text, font=FONT_CHIP, fill="#111827")


PALETTES = {
    "blue": {"fill": "#EFF6FF", "outline": "#60A5FA", "accent": "#2563EB"},
    "green": {"fill": "#ECFDF5", "outline": "#4ADE80", "accent": "#16A34A"},
    "orange": {"fill": "#FFF7ED", "outline": "#FB923C", "accent": "#EA580C"},
    "purple": {"fill": "#F5F3FF", "outline": "#A78BFA", "accent": "#7C3AED"},
    "pink": {"fill": "#FDF2F8", "outline": "#F472B6", "accent": "#DB2777"},
    "slate": {"fill": "#F8FAFC", "outline": "#CBD5E1", "accent": "#64748B"},
}


STAGES = [
    {
        "file": "01_document_parsing_workcard.png",
        "title": "01 文档解析",
        "summary": "把多格式原始资料变成可追踪、可分块、可入库的标准化文本与结构化对象。",
        "color": "blue",
        "inputs": ["PDF / Word / TXT / JSON / 网页", "扫描件、图片、工程图纸", "文件来源、版本、权限、上传时间"],
        "actions": ["文件登记", "解析器选择", "版面恢复", "OCR 校正", "页码与 metadata 绑定"],
        "middle": ["原件 checksum", "pages / blocks", "tables / images", "OCR 置信度日志"],
        "outputs": ["标准化解析包", "可分块文本", "低置信区域复核清单"],
        "quality": ["页码和原文可回溯", "表格不丢列", "低质量 OCR 显式标记"],
    },
    {
        "file": "02_entity_relation_extraction_workcard.png",
        "title": "02 实体关系抽取",
        "summary": "从文本片段中抽取设备、部件、故障、原因、措施等实体和带证据的关系。",
        "color": "green",
        "inputs": ["解析文本和 chunk", "专业 schema / 术语词表", "故障模式模板和人工规则"],
        "actions": ["候选实体识别", "关系候选生成", "同义词归并", "Evidence 绑定", "人工审核"],
        "middle": ["entity_candidates", "relation_candidates", "triples_draft", "review_decisions"],
        "outputs": ["entities.json", "relations.json", "triples.json", "evidence_links.json"],
        "quality": ["无 evidence 不入库", "实体类型受 schema 约束", "关系方向和语义可解释"],
    },
    {
        "file": "03_graph_construction_workcard.png",
        "title": "03 图谱构建",
        "summary": "把审核后的实体与关系组织成可检索、可计算、可版本管理的知识图谱。",
        "color": "orange",
        "inputs": ["审核后的实体和三元组", "证据链接", "图谱 schema 和构建配置"],
        "actions": ["实体对齐", "关系归一", "属性补全", "图结构校验", "版本发布"],
        "middle": ["nodes / edges 草稿", "alias_map", "graph_stats", "build_diff"],
        "outputs": ["GraphStore", "nodes.csv", "edges.csv", "graph_manifest"],
        "quality": ["节点边可回到证据", "非法关系被拦截", "每次构建可复现"],
    },
    {
        "file": "04_community_detection_workcard.png",
        "title": "04 社区划分",
        "summary": "把大图按结构关联和业务主题拆成更容易解释、摘要和调用的知识社区。",
        "color": "purple",
        "inputs": ["图谱节点与边", "边权和关系类型", "专业层级与人工边界"],
        "actions": ["图投影与加权", "算法划分", "子图抽取", "社区解释", "人工边界调整"],
        "middle": ["community_map", "subgraphs", "central_nodes", "cross_edges"],
        "outputs": ["社区编号", "社区子图", "核心节点列表", "划分日志"],
        "quality": ["社区有业务含义", "不过度碎片化", "跨社区关系保留"],
    },
    {
        "file": "05_topic_summary_workcard.png",
        "title": "05 主题摘要",
        "summary": "为每个知识社区生成可引用、可复核、可进入 GraphRAG 上下文的主题说明。",
        "color": "pink",
        "inputs": ["社区子图", "核心节点和关键路径", "代表 evidence", "历史摘要"],
        "actions": ["代表证据选择", "主题归纳", "标签化", "冲突检查", "专家复核"],
        "middle": ["summary_draft", "summary_evidence", "topic_tags", "conflict_notes"],
        "outputs": ["community_summary", "topic_tags", "global_context"],
        "quality": ["摘要带代表证据", "不编造关系", "冲突内容显式标注"],
    },
    {
        "file": "06_evidence_retrieval_workcard.png",
        "title": "06 证据召回",
        "summary": "围绕用户问题从文本、图谱路径和社区摘要中召回可支撑回答的证据包。",
        "color": "blue",
        "inputs": ["用户问题和历史上下文", "向量库 / 关键词索引", "图谱路径和社区摘要"],
        "actions": ["问题理解", "查询改写", "多路召回", "重排去重", "证据包封装"],
        "middle": ["query_rewrite", "retrieval_candidates", "rerank_scores", "graph_paths"],
        "outputs": ["TopK chunk", "图谱路径", "社区摘要", "evidence_bundle"],
        "quality": ["关键实体被覆盖", "来源和页码齐全", "重复证据被合并"],
    },
    {
        "file": "07_reasoning_generation_workcard.png",
        "title": "07 推理生成",
        "summary": "基于证据包和图谱上下文生成有引用、有边界、有结构的回答或报告。",
        "color": "green",
        "inputs": ["用户问题", "evidence_bundle", "图谱路径 / 社区摘要", "回答模板"],
        "actions": ["Prompt 装配", "依据链组织", "答案生成", "引用核对", "结构化落盘"],
        "middle": ["prompt_snapshot", "reasoning_basis", "citation_map", "uncertainty_flags"],
        "outputs": ["answer.md", "answer.json", "report.docx", "引用映射"],
        "quality": ["无证据则声明不确定", "关键结论有引用", "输出格式稳定"],
    },
    {
        "file": "08_result_evaluation_workcard.png",
        "title": "08 结果评估",
        "summary": "对召回、生成、引用和图谱增强效果做评估，并把失败案例转成下一轮改进任务。",
        "color": "slate",
        "inputs": ["问题集和标准答案", "模型回答", "证据包", "运行日志"],
        "actions": ["样本归档", "指标统计", "质量评审", "失败归因", "改进闭环"],
        "middle": ["eval_cases", "metrics", "judge_notes", "failure_cases"],
        "outputs": ["evaluation_report", "metrics.csv", "failure_cases", "improvement_backlog"],
        "quality": ["样本可复现", "失败可定位", "指标支撑下一轮优化"],
    },
]


def draw_stage(stage: dict[str, object]) -> Path:
    width, height = 2400, 1500
    image = Image.new("RGB", (width, height), "#F8FAFC")
    draw = ImageDraw.Draw(image)
    palette = PALETTES[str(stage["color"])]

    draw.text((76, 50), str(stage["title"]), font=FONT_TITLE, fill="#0F172A")
    draw_wrapped(draw, (80, 115), str(stage["summary"]), FONT_SUBTITLE, "#475569", 1500, line_gap=34)
    draw.rounded_rectangle((78, 178, 2320, 188), radius=5, fill=palette["accent"])
    draw.text((1920, 88), "工作卡版：不限定步骤数", font=FONT_SMALL, fill="#64748B")

    input_rect = (80, 260, 540, 670)
    core_rect = (610, 245, 1790, 710)
    output_rect = (1860, 260, 2320, 670)
    middle_rect = (80, 790, 1135, 1230)
    quality_rect = (1225, 790, 2320, 1230)
    keep_rect = (80, 1280, 2320, 1410)

    panel(draw, input_rect, "输入与触发", palette)
    bullet_list(draw, list(stage["inputs"]), input_rect[0] + 36, input_rect[1] + 92, 370)

    rounded(draw, core_rect, "#FFFFFF", palette["outline"], width=4, radius=28)
    draw.text((core_rect[0] + 36, core_rect[1] + 28), "核心动作池", font=FONT_H2, fill="#0F172A")
    draw.text((core_rect[0] + 200, core_rect[1] + 32), "可按资料类型和任务目标增删组合", font=FONT_SMALL, fill="#64748B")
    actions = list(stage["actions"])
    cols = 3
    chip_w = 330
    chip_h = 92
    gap_x = 40
    gap_y = 36
    start_x = core_rect[0] + 66
    start_y = core_rect[1] + 115
    for idx, action in enumerate(actions):
        row = idx // cols
        col = idx % cols
        x = start_x + col * (chip_w + gap_x)
        y = start_y + row * (chip_h + gap_y)
        chip(draw, (x, y, x + chip_w, y + chip_h), str(action), palette)
    draw.rounded_rectangle((core_rect[0] + 66, core_rect[3] - 102, core_rect[2] - 66, core_rect[3] - 34), radius=18, fill="#FFFFFF", outline="#CBD5E1", width=2)
    draw.text((core_rect[0] + 94, core_rect[3] - 84), "说明：这里表达“能力集合”，不是固定流水线；具体项目可按数据质量、专业要求和时间成本裁剪。", font=FONT_SMALL, fill="#475569")

    panel(draw, output_rect, "出口与交付", palette)
    bullet_list(draw, list(stage["outputs"]), output_rect[0] + 36, output_rect[1] + 92, 370)

    arrow(draw, (input_rect[2], 465), (core_rect[0], 465), "#64748B", 5)
    arrow(draw, (core_rect[2], 465), (output_rect[0], 465), "#64748B", 5)

    panel(draw, middle_rect, "中间结果留存", palette)
    mids = list(stage["middle"])
    x_positions = [middle_rect[0] + 42, middle_rect[0] + 540]
    y_positions = [middle_rect[1] + 105, middle_rect[1] + 245]
    for idx, item in enumerate(mids):
        x = x_positions[idx % 2]
        y = y_positions[idx // 2]
        rounded(draw, (x, y, x + 425, y + 95), palette["fill"], palette["outline"], width=2, radius=16)
        draw.text((x + 24, y + 30), str(item), font=FONT_BODY, fill="#111827")

    panel(draw, quality_rect, "质量判断与复核点", palette)
    bullet_list(draw, list(stage["quality"]), quality_rect[0] + 42, quality_rect[1] + 105, 950, max_lines_each=2)
    rounded(draw, (quality_rect[0] + 42, quality_rect[3] - 112, quality_rect[2] - 42, quality_rect[3] - 36), "#FFFFFF", "#CBD5E1", width=2, radius=18)
    draw.text((quality_rect[0] + 70, quality_rect[3] - 88), "建议沉淀：失败样本、人工修改意见、参数配置和下一轮优化任务。", font=FONT_SMALL, fill="#475569")

    rounded(draw, keep_rect, "#FFFFFF", palette["outline"], width=2, radius=24)
    draw.text((keep_rect[0] + 34, keep_rect[1] + 36), "统一留存口径", font=FONT_CARD_TITLE, fill="#0F172A")
    draw.text(
        (keep_rect[0] + 210, keep_rect[1] + 39),
        "每个产物文件建议包含 stage、run_id、数据版本、模型/参数、时间戳；关键结论必须能回到原文 evidence。",
        font=FONT_BODY,
        fill="#334155",
    )

    out = OUT_DIR / str(stage["file"])
    out.parent.mkdir(parents=True, exist_ok=True)
    image.save(out, quality=95)
    return out


def set_run_font(run, size: int | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color.replace("#", ""))


def add_paragraph(doc: Document, text: str, size: int = 10, bold: bool = False, color: str = "#111827", after: int = 6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def build_docx(image_paths: list[Path]) -> None:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.top_margin = Inches(0.35)
    section.bottom_margin = Inches(0.35)
    section.left_margin = Inches(0.35)
    section.right_margin = Inches(0.35)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("PowerRAG / GraphRAG 关键环节流程图合集：工作卡版")
    set_run_font(run, size=23, bold=True, color="#0F172A")
    add_paragraph(doc, "日期：2026-07-07    说明：本版不把每个环节固化成相同数量的流水步骤，而是用“输入、动作池、留存、出口、质量判断”表达可裁剪的工作方式。", size=10, color="#475569")
    add_paragraph(doc, "适合放进 WPS 过程结果文档，用来说明每个环节要沉淀哪些中间结果和最终输出。", size=11, color="#111827", after=10)

    for idx, stage in enumerate(STAGES, start=1):
        add_paragraph(doc, f"{idx}. {stage['title']}", size=12, bold=True, color="#0F172A", after=2)
        add_paragraph(doc, str(stage["summary"]), size=10, color="#475569", after=5)
    doc.add_page_break()

    for idx, (stage, img_path) in enumerate(zip(STAGES, image_paths), start=1):
        h = doc.add_paragraph()
        h.paragraph_format.space_after = Pt(4)
        run = h.add_run(str(stage["title"]))
        set_run_font(run, size=16, bold=True, color="#0F172A")
        desc = doc.add_paragraph()
        desc.paragraph_format.space_after = Pt(6)
        r = desc.add_run(str(stage["summary"]))
        set_run_font(r, size=9, color="#475569")

        pic_p = doc.add_paragraph()
        pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic_run = pic_p.add_run()
        pic_run.add_picture(str(img_path), width=Inches(10.55))
        inline = doc.inline_shapes[-1]._inline
        inline.docPr.set("descr", str(stage["title"]))
        inline.docPr.set("title", str(stage["title"]))

        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(4)
        cr = cap.add_run(f"图 {idx}：{stage['title']}工作卡")
        set_run_font(cr, size=9, color="#475569")
        if idx != len(STAGES):
            doc.add_page_break()

    doc.core_properties.title = "PowerRAG / GraphRAG 关键环节流程图合集：工作卡版"
    doc.core_properties.subject = "PowerRAG flexible workflow cards"
    doc.core_properties.author = "Codex"
    doc.core_properties.keywords = "PowerRAG, GraphRAG, workflow card, RAG"
    doc.core_properties.comments = ""
    doc.save(DOCX_OUT)


def build_markdown(image_paths: list[Path]) -> None:
    lines = [
        "# PowerRAG / GraphRAG 关键环节流程图合集：工作卡版",
        "",
        "日期：2026-07-07",
        "",
        "本版不再把每个环节固定画成同样数量的流程步骤，而是用“输入与触发、核心动作池、中间结果留存、出口与交付、质量判断与复核点”表达可裁剪的工作方式。",
        "",
    ]
    for idx, (stage, img_path) in enumerate(zip(STAGES, image_paths), start=1):
        rel = img_path.relative_to(MD_OUT.parent).as_posix()
        lines.extend(
            [
                f"## {idx}. {stage['title']}",
                "",
                str(stage["summary"]),
                "",
                f"![{stage['title']}]({rel})",
                "",
            ]
        )
    MD_OUT.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    image_paths = [draw_stage(stage) for stage in STAGES]
    build_docx(image_paths)
    build_markdown(image_paths)
    print(DOCX_OUT)
    print(MD_OUT)
    for path in image_paths:
        print(path)


if __name__ == "__main__":
    main()
