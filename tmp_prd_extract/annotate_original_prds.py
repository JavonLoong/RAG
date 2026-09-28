# -*- coding: utf-8 -*-
"""Copy original PRDs and prepend ☑ only. Do not rewrite any narrative."""

from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

MARK = "☑ "

MAIN_SRC = Path(r"D:\虚拟C盘\PowerRAG PRD.docx")
MAIN_DST = Path(r"D:\虚拟C盘\PowerRAG PRD_勾选核对.docx")
V10_SRC = Path(r"D:\虚拟C盘\PowerRAG-latest\docs\PowerRAG_GraphRAG_Governance_Workbench_PRD_v1.0_2026-08-09.docx")
V10_DST = Path(r"D:\虚拟C盘\PowerRAG_工作台PRD_v1.0_勾选核对.docx")

# Only items treated as done (code wired into the product). Partial / not accepted stay unmarked.
MAIN_PREFIXES = (
    "PR-P0-01 ",
    "PR-P0-02 ",
    "PR-P0-03 ",
    "PR-P0-04 ",
    "PR-P0-05 ",
    "PR-P0-06 ",
    "PR-P0-07 ",
    "PR-P0-08 ",
    "PR-P0-09 ",
    "PR-P0-10 ",
    "PR-P0-11 ",
    "PR-P0-13 ",
    "PR-P0-14 ",
    "PR-P0-15 ",
    "PR-P0-16 ",
    "PR-P0-17 ",
    "PR-P0-18 ",
    "PR-P0-19 ",
    "PR-P0-20 ",
    "PR-P0-21 ",
    "PR-P0-22 ",
    "PR-P0-23 ",
    "PR-P0-24 ",
    "PR-P0-25 ",
    "PR-P0-26 ",
    "PR-P0-27 ",
    "PR-P1-01 ",
    "PR-P1-03 ",
    "TECH-01 ",
    "TECH-02 ",
    "TECH-03 ",
    "TECH-04 ",
    "TECH-05 ",
    "TECH-06 ",
    "TECH-07 ",
    "TECH-08 ",
    "TECH-09 ",
    "TECH-10 ",
    "TECH-11 ",
    "TECH-12 ",
    "TECH-13 ",
    "TECH-14 ",
    "TECH-16 ",
    "（一）建立项目与数据基线",
    "（二）资料接入、解析和质检",
    "（三）正式资料发布与检索",
    "（六）FMEA 生成与交付",
    "（七）问题反馈与重验证",
    "（二）阶段 1：产品主链集成",
)

V10_CELL_PREFIXES = (
    "FR-001",
    "FR-002",
    "FR-003",
    "FR-004",
    "FR-005",
    "FR-006",
    "FR-007",
    "FR-008",
    "FR-009",
    "FR-010",
    "FR-012",
    "FR-013",
    "FR-014",
    "FR-015",
    "FR-016",
    "FR-017",
    "FR-018",
    "FR-019",
    "AC-01",
    "AC-02",
    "AC-03",
    "AC-04",
    "AC-05",
    "AC-06",
    "AC-07",
    "AC-08",
    "AC-09",
    "AC-10",
    "AC-11",
    "NFR-02",
    "NFR-03",
    "NFR-04",
    "NFR-05",
    "NFR-07",
    "GET /api/delivery/review-queue",
    "GET /api/delivery/documents",
    "GET /api/delivery/fmea/tasks",
    "GET /api/delivery/audit/events",
    "GET /api/delivery/batches/{id}",
    "发布接口请求体",
    "统一错误结构",
    "交付总览",
    "资料审核",
    "图谱审核",
    "FMEA 交付",
    "审计与反馈",
    "系统配置",
    "阶段 1：资料审核 MVP",
    "阶段 2：图谱与 FMEA",
)


def prepend_first_paragraph(paragraph, mark=MARK):
    text = paragraph.text or ""
    if text.startswith(mark):
        return False
    if not paragraph.runs:
        run = paragraph.add_run(mark)
        return True
    first = paragraph.runs[0]._r
    new_r = OxmlElement("w:r")
    rpr = first.find(qn("w:rPr"))
    if rpr is not None:
        new_r.append(deepcopy(rpr))
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = mark
    new_r.append(t)
    first.addprevious(new_r)
    return True


def starts_with_any(text, prefixes):
    return any(text.startswith(p) for p in prefixes)


def annotate_main(path):
    doc = Document(path)
    n = 0
    for para in doc.paragraphs:
        text = para.text or ""
        if starts_with_any(text, MAIN_PREFIXES):
            if prepend_first_paragraph(para):
                n += 1
    doc.save(path)
    return n


def annotate_v10(path):
    doc = Document(path)
    n = 0
    for table in doc.tables:
        for row in table.rows:
            cell = row.cells[0]
            text = (cell.text or "").strip()
            if not starts_with_any(text, V10_CELL_PREFIXES):
                continue
            if cell.paragraphs and prepend_first_paragraph(cell.paragraphs[0]):
                n += 1
    # Journey table first-column step numbers 1-7 (already-built UI chain)
    journey = None
    for table in doc.tables:
        if len(table.columns) == 5 and table.rows and table.rows[0].cells[0].text.strip() == "步骤":
            journey = table
            break
    if journey is not None:
        for row in journey.rows[1:]:
            step = (row.cells[0].text or "").strip()
            if step in {"1", "2", "3", "4", "5", "6", "7"}:
                if prepend_first_paragraph(row.cells[0].paragraphs[0]):
                    n += 1
    doc.save(path)
    return n


def main():
    for src, dst in ((MAIN_SRC, MAIN_DST), (V10_SRC, V10_DST)):
        shutil.copy2(src, dst)
    a = annotate_main(MAIN_DST)
    b = annotate_v10(V10_DST)
    print(f"main marks={a} {MAIN_DST} {MAIN_DST.stat().st_size}")
    print(f"v10 marks={b} {V10_DST} {V10_DST.stat().st_size}")


if __name__ == "__main__":
    main()
