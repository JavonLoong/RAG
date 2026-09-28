# -*- coding: utf-8 -*-
"""按口述改虚拟 C 盘展示稿：中选可见、去掉测试项、一条线改优势表述。"""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

SRC = Path(r"D:\虚拟C盘\纪文龙_M2-M5工作展示_20260923.pptx")
COPIES = [
    SRC,
    Path(r"D:\虚拟C盘\PowerRAG-latest\docs\project_deliverables\06_汇报材料_发群和组会\纪文龙_M2-M5工作展示_20260923.pptx"),
]

BLACK = RGBColor(0x00, 0x00, 0x00)
FONT = "微软雅黑"
FOOT_TOP = 5989320
FOOT_H = 620000


def _font(run, size: int) -> None:
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    run.font.bold = False
    run.font.name = FONT
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", FONT)


def set_box(shape, value: str, size: int = 22) -> None:
    tf = shape.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, part in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = part
        _font(run, size)


def set_footer(shape, value: str, size: int = 13) -> None:
    shape.top = Emu(FOOT_TOP)
    shape.height = Emu(FOOT_H)
    set_box(shape, value, size)


def add_footer(slide, value: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.45), Emu(FOOT_TOP), Inches(12.4), Emu(FOOT_H))
    set_footer(box, value)


def is_footer(shape) -> bool:
    return shape.has_text_frame and int(shape.top) >= 5_700_000


def set_cell(cell, value: str, size: int = 14, bold: bool = False) -> None:
    tf = cell.text_frame
    tf.word_wrap = True
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = value
    _font(run, size)
    run.font.bold = bold
    body = tf._txBody.find(qn("a:bodyPr"))
    if body is not None:
        body.set("anchor", "ctr")


def delete_last_row(table) -> None:
    tbl = table._tbl
    rows = list(tbl.tr_lst)
    if len(rows) < 2:
        return
    tbl.remove(rows[-1])


def patch(prs: Presentation) -> None:
    s1, s2, s3, s4, s5 = prs.slides

    # 封面：M2 到 M5 怎么串
    set_box(
        s1.shapes[1],
        "M2 解析完，交给 M3 入库。\nM4 只吃已入库资料构图。M5 只吃已发布图出表。",
        22,
    )
    for shape in s1.shapes:
        if is_footer(shape):
            set_footer(shape, "2026-09-23    FMEA：故障模式与影响分析。出表字段必须能回到原文页码。")

    # 分工 → 分别作用，去掉测试项
    set_box(s2.shapes[0], "分别作用", 22)
    set_box(
        s2.shapes[1],
        "M2 解析：把文件读成带页码的文字和结构。\n"
        "M3 入库：审过的资料做成可检索的正式版本。\n"
        "M4 图谱：从正式资料抽出设备、故障、原因、措施，每条绑回原页。\n"
        "M5 输出：从正式图生成可审核的 FMEA 表。",
        22,
    )
    found_footer = False
    for shape in s2.shapes:
        if is_footer(shape):
            set_footer(shape, "FMEA：故障模式与影响分析。后一段只用前一段已经发布的结果。")
            found_footer = True
    if not found_footer:
        add_footer(s2, "FMEA：故障模式与影响分析。后一段只用前一段已经发布的结果。")

    # 一条线：删人审句，写清相对直接丢给大模型的优势
    set_box(
        s3.shapes[1],
        "资料进来，读成带页码的证据。\n"
        "只从已发布资料构图，关系必须绑回原页。\n"
        "这是相对把资料直接丢给大模型作答的优势。\n"
        "只从已发布图出 FMEA，每个字段必须带证据。",
        22,
    )
    foot3 = "解析  —  入库  —  图谱  —  FMEA\n绑证据：每条结论能回到手册页码，不是模型随口编的。"
    found_footer = False
    for shape in s3.shapes:
        if is_footer(shape):
            set_footer(shape, foot3)
            found_footer = True
    if not found_footer:
        add_footer(s3, foot3)

    # 做成了什么：标红术语改白话
    table = next(shape.table for shape in s4.shapes if shape.has_table)
    set_cell(table.cell(1, 1), "识图取字，保住页码结构，对照原页审", 14)
    set_cell(table.cell(2, 2), "前 5 条检索都能对上；四类边界过；引用可回原块", 14)
    set_cell(table.cell(3, 2), "审核过 120 条关系；发布图 18 条都能回到原页；同一批问题对照过 4 道", 14)
    set_cell(table.cell(4, 1), "已发布图加模板，字段带证据，再导出", 14)
    set_cell(table.cell(4, 2), "3 条专家审过；JSON / CSV / DOCX 行数一致", 14)
    set_cell(table.cell(5, 0), "工作台", 14, bold=True)
    set_cell(table.cell(5, 1), "在工作台上走完", 14)
    set_cell(table.cell(5, 2), "试点项目已建；资料、图、FMEA 已发布；回滚和恢复已跑通", 14)
    set_cell(table.cell(6, 1), "工程项通过才往下交", 14)
    set_cell(table.cell(6, 2), "工程 8 项通过；端到端还差领域审核签字", 14)
    foot4 = "前 5 条：检索结果最靠前的 5 段。门禁：验收门槛。FMEA：故障模式与影响分析。"
    found_footer = False
    for shape in s4.shapes:
        if is_footer(shape):
            set_footer(shape, foot4)
            found_footer = True
    if not found_footer:
        add_footer(s4, foot4)

    # 门禁表：去掉测试负责人行，条款号改成模块名
    gate = next(shape.table for shape in s5.shapes if shape.has_table)
    labels = ["解析", "入库检索", "图谱", "FMEA 导出", "端到端"]
    for i, label in enumerate(labels, start=1):
        set_cell(gate.cell(i, 0), label, 16, bold=True)
    set_cell(gate.cell(2, 2), "前 5 条检索都能对上，四类边界过", 16)
    set_cell(gate.cell(5, 2), "差领域审核签字。必须两个角色只点界面", 16)
    last_label = gate.cell(len(gate.rows) - 1, 0).text_frame.text
    if "TECH" in last_label or "测试" in last_label:
        delete_last_row(gate)
    found_footer = False
    for shape in s5.shapes:
        if is_footer(shape):
            set_footer(shape, "门禁：验收门槛。端到端：从建项走到导出和恢复。")
            found_footer = True
    if not found_footer:
        add_footer(s5, "门禁：验收门槛。端到端：从建项走到导出和恢复。")


def main() -> None:
    prs = Presentation(str(SRC))
    if len(prs.slides) != 5:
        raise SystemExit(f"expected 5 slides, got {len(prs.slides)}")
    patch(prs)
    for dest in COPIES:
        dest.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(dest))
        print(dest)


if __name__ == "__main__":
    main()
