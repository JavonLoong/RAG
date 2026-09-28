# -*- coding: utf-8 -*-
from pathlib import Path
from pptx import Presentation

SRC = Path(r"D:\虚拟C盘\纪文龙_M2-M5工作展示_20260923.pptx")
out = Path(r"D:\虚拟C盘\PowerRAG-latest\_ppt_now.txt")
prs = Presentation(str(SRC))
lines = [f"slides={len(prs.slides)}"]
for i, slide in enumerate(prs.slides, 1):
    lines.append(f"\n==== SLIDE {i}")
    for j, shape in enumerate(slide.shapes):
        if shape.has_table:
            lines.append(f"  TBL{j} rows={len(shape.table.rows)} cols={len(shape.table.columns)}")
            for r, row in enumerate(shape.table.rows):
                cells = [c.text_frame.text.replace("\n", " / ") for c in row.cells]
                lines.append(f"    r{r}: {' | '.join(cells)}")
        elif shape.has_text_frame:
            text = shape.text_frame.text.replace("\n", " / ")
            lines.append(f"  T{j} top={int(shape.top)} h={int(shape.height)} {text}")
out.write_text("\n".join(lines), encoding="utf-8")
print(out)
