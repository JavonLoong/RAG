# -*- coding: utf-8 -*-
from pathlib import Path
from pptx import Presentation

src = Path(r"D:\虚拟C盘\PowerRAG_整合版_技术链路与接口_20260923.pptx")
prs = Presentation(str(src))
out = []
out.append(f"slides={len(prs.slides)}")
for i, slide in enumerate(prs.slides, 1):
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip().split("\n")[0]
            if t:
                texts.append(t)
                break
        if shape.has_table:
            texts.append(f"table {len(shape.table.rows)}x{len(shape.table.columns)}")
            break
    out.append(f"{i:02d} {texts[0] if texts else '?'}")
Path(r"D:\虚拟C盘\PowerRAG-latest\_ppt_full_now.txt").write_text("\n".join(out), encoding="utf-8")
print("\n".join(out))
